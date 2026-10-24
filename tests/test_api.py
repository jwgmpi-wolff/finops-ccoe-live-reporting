import asyncio
import subprocess
from pathlib import Path
from fastapi.testclient import TestClient
from datetime import UTC, datetime
import json
import tempfile
import pytest
import app.main as main_module
from pptx import Presentation
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Inches

from app.main import LAST_REPORT, app
from app.models import Recommendation, ReportDomain, ReportResult, SourceAudit
from app.services.exporters import build_executive_brief, export_pdf, export_pptx
from app.services.diagrams import build_report_diagram
from app.services.reporting import build_report
from app.services.azure_live import AzureLiveClient


client = TestClient(app)


def test_health_reports_deployment_version():
    response = client.get("/api/health")

    assert response.status_code == 200
    assert response.json()["version"] == "1.2.5"


def test_tenant_spend_columns_open_meter_level_drilldown():
    app_script = (Path(__file__).parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")

    assert "function openSpendDrilldown(subscriptionId, usageDate, rows)" in app_script
    assert 'item.setAttribute("role", "button")' in app_script
    assert 'event.key === "Enter" || event.key === " "' in app_script
    assert '"meterSubCategory"' in app_script
    assert '"unitOfMeasure"' in app_script
    assert "select a bar for meter-level detail" in app_script


def test_browser_visual_contract_covers_every_report_domain():
    app_script = (Path(__file__).parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")
    domain_block = app_script.split("const visualReportDomains = new Set([", 1)[1].split("]);", 1)[0]

    for domain in ReportDomain:
        assert f'"{domain.value}"' in domain_block, domain.value


def test_cli_login_uses_hidden_browser_flow(monkeypatch):
    captured = {}

    class FakeProcess:
        def poll(self):
            return None

    def fake_popen(command, **kwargs):
        captured["command"] = command
        captured["kwargs"] = kwargs
        return FakeProcess()

    monkeypatch.setattr("app.main._resolve_az_command", lambda: "az.cmd")
    monkeypatch.setattr("app.main.CLI_LOGIN_PROCESS", None)
    monkeypatch.setattr("app.main.CLI_LOGIN_OUTPUT", None)
    monkeypatch.setattr("app.main.subprocess.Popen", fake_popen)

    response = client.post("/api/auth/cli-login")

    assert response.status_code == 200
    assert response.json()["status"] == "started"
    assert captured["kwargs"]["stdin"] is subprocess.DEVNULL
    assert captured["kwargs"]["stdout"] is main_module.CLI_LOGIN_OUTPUT
    assert captured["kwargs"]["stderr"] is subprocess.STDOUT
    assert "--use-device-code" not in captured["command"]
    assert "--tenant" not in captured["command"]
    assert captured["kwargs"]["env"]["AZURE_CORE_ENABLE_BROKER_ON_WINDOWS"] == "false"
    assert captured["kwargs"]["env"]["AZURE_CORE_LOGIN_EXPERIENCE_V2"] == "off"
    main_module._close_cli_login_output()


def test_cli_device_login_supports_external_tenant(monkeypatch):
    captured = {}

    class FakeProcess:
        def poll(self):
            return None

    def fake_popen(command, **kwargs):
        captured["command"] = command
        return FakeProcess()

    monkeypatch.setattr(main_module, "_resolve_az_command", lambda: "az")
    monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", None)
    monkeypatch.setattr(main_module, "CLI_LOGIN_OUTPUT", None)
    monkeypatch.setattr(main_module.subprocess, "Popen", fake_popen)

    response = client.post("/api/auth/cli-login?use_device_code=true&tenant_id=contoso.onmicrosoft.com")

    assert response.status_code == 200
    assert captured["command"] == [
        "az", "login", "--allow-no-subscriptions", "--tenant",
        "contoso.onmicrosoft.com", "--output", "none", "--use-device-code",
    ]
    main_module._close_cli_login_output()


@pytest.mark.parametrize("return_code, expected_status", [(None, "pending"), (0, "completed"), (1, None)])
def test_cli_login_status_surfaces_challenge_completion_and_failure(monkeypatch, return_code, expected_status):
    class FakeProcess:
        def poll(self):
            return return_code

    with tempfile.NamedTemporaryFile(mode="w+t", encoding="utf-8", delete=False) as output:
        output.write(
            "To sign in, use a web browser to open https://microsoft.com/devicelogin "
            "and enter the code ABC123XYZ to authenticate.\n"
            "AADSTS53003: Your organization blocked this sign-in.\n"
        )
        output.flush()
        monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", FakeProcess())
        monkeypatch.setattr(main_module, "CLI_LOGIN_OUTPUT", output)
        try:
            position = output.tell()
            response = client.get("/api/auth/cli-login-status")
            assert output.tell() == position
            if return_code == 1:
                assert response.status_code == 502
                assert "AADSTS53003" in response.json()["detail"]
            else:
                assert response.status_code == 200
                assert response.json()["status"] == expected_status
                if return_code is None:
                    assert response.json()["user_code"] == "ABC123XYZ"
                    assert response.json()["verification_uri"] == "https://microsoft.com/devicelogin"
                else:
                    assert "user_code" not in response.json()
        finally:
            main_module._close_cli_login_output()


def test_cli_login_missing_cli_reports_prerequisite(monkeypatch):
    monkeypatch.setattr(main_module, "which", lambda _: None)
    monkeypatch.setattr(main_module.FsPath, "exists", lambda _: False)

    response = client.post("/api/auth/cli-login?use_device_code=true")

    assert response.status_code == 502
    assert "Azure CLI not found" in response.json()["detail"]


def test_installed_app_uses_bundled_cli_without_path(monkeypatch, tmp_path):
    executable = tmp_path / "Programs With Spaces" / "FinOpsCCoE.exe"
    bundled_cli = executable.parent / "azure-cli" / "bin" / "az.cmd"
    bundled_cli.parent.mkdir(parents=True)
    bundled_cli.touch()
    monkeypatch.setattr(main_module.sys, "frozen", True, raising=False)
    monkeypatch.setattr(main_module.sys, "executable", str(executable))

    def no_system_search(_):
        raise AssertionError("Bundled sign-in must not depend on a system CLI or PATH")

    monkeypatch.setattr(main_module, "which", no_system_search)

    assert main_module._resolve_az_command() == str(bundled_cli)


def test_installed_app_missing_runtime_requests_repair(monkeypatch, tmp_path):
    monkeypatch.setattr(main_module.sys, "frozen", True, raising=False)
    monkeypatch.setattr(main_module.sys, "executable", str(tmp_path / "FinOpsCCoE.exe"))
    monkeypatch.setattr(main_module, "which", lambda _: None)
    monkeypatch.setattr(main_module.FsPath, "exists", lambda _: False)

    response = client.post("/api/auth/cli-login")

    assert response.status_code == 502
    assert "Reinstall the latest Windows package" in response.json()["detail"]


def test_cli_token_uses_same_authentication_environment(monkeypatch):
    captured = {}

    def fake_run(command, **kwargs):
        captured.update(kwargs)
        return subprocess.CompletedProcess(command, 0, "test-token\n", "")

    monkeypatch.setattr(main_module, "_resolve_az_command", lambda: "az")
    monkeypatch.setattr(main_module.subprocess, "run", fake_run)

    response = client.get("/api/auth/cli-token")

    assert response.status_code == 200
    assert response.json()["access_token"] == "test-token"
    assert captured["env"]["AZURE_CORE_ENABLE_BROKER_ON_WINDOWS"] == "false"
    assert captured["env"]["AZURE_CORE_LOGIN_EXPERIENCE_V2"] == "off"


def test_installed_authentication_isolates_profile_and_extensions(monkeypatch, tmp_path):
    monkeypatch.setattr(main_module.sys, "frozen", True, raising=False)
    monkeypatch.setenv("LOCALAPPDATA", str(tmp_path))
    monkeypatch.setenv("AZURE_CONFIG_DIR", str(tmp_path / "system-cli"))
    monkeypatch.setenv("AZURE_EXTENSION_DIR", str(tmp_path / "system-extensions"))

    environment = main_module._azure_cli_login_environment()

    assert environment["AZURE_CONFIG_DIR"] == str(tmp_path / "FinOpsCCoE" / "azure")
    assert environment["AZURE_EXTENSION_DIR"] == str(tmp_path / "FinOpsCCoE" / "azure" / "cliextensions")


def test_cli_login_rejects_invalid_tenant():
    response = client.post("/api/auth/cli-login?tenant_id=contoso%26other")

    assert response.status_code == 422


def test_cli_login_launch_failure_cleans_up_output(monkeypatch):
    def fail_to_start(*args, **kwargs):
        raise OSError("Azure CLI cannot start")

    monkeypatch.setattr(main_module, "_resolve_az_command", lambda: "az")
    monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", None)
    monkeypatch.setattr(main_module, "CLI_LOGIN_OUTPUT", None)
    monkeypatch.setattr(main_module.subprocess, "Popen", fail_to_start)

    response = client.post("/api/auth/cli-login")

    assert response.status_code == 502
    assert "Azure CLI cannot start" in response.json()["detail"]
    assert main_module.CLI_LOGIN_OUTPUT is None


def test_cli_login_does_not_hide_session_reset_failure(monkeypatch):
    monkeypatch.setattr(main_module, "_resolve_az_command", lambda: "az")
    monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", None)
    monkeypatch.setattr(
        main_module.subprocess, "run",
        lambda *args, **kwargs: subprocess.CompletedProcess(args[0], 1, "", "CLI session reset denied"),
    )

    response = client.post("/api/auth/cli-login?force_reauth=true")

    assert response.status_code == 502
    assert "CLI session reset denied" in response.json()["detail"]


@pytest.mark.parametrize("force_reauth", [False, True])
def test_cli_login_starts_without_existing_accounts(monkeypatch, force_reauth):
    commands = []

    class FakeProcess:
        def poll(self):
            return None

    def fake_run(command, **kwargs):
        commands.append(command)
        if command == ["az", "logout"]:
            return subprocess.CompletedProcess(command, 1, "", "ERROR: There are no active accounts.\n")
        return subprocess.CompletedProcess(command, 0, "", "")

    def fake_popen(command, **kwargs):
        commands.append(command)
        return FakeProcess()

    monkeypatch.setattr(main_module, "_resolve_az_command", lambda: "az")
    monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", None)
    monkeypatch.setattr(main_module, "CLI_LOGIN_OUTPUT", None)
    monkeypatch.setattr(main_module.subprocess, "run", fake_run)
    monkeypatch.setattr(main_module.subprocess, "Popen", fake_popen)

    response = client.post(
        f"/api/auth/cli-login?force_reauth={str(force_reauth).lower()}&use_device_code=true"
    )

    try:
        assert response.status_code == 200
        assert response.json()["status"] == "started"
        assert commands[-1] == ["az", "login", "--allow-no-subscriptions", "--output", "none", "--use-device-code"]
        assert len(commands) == (3 if force_reauth else 1)
    finally:
        main_module._close_cli_login_output()


def test_browser_sign_in_does_not_force_global_cli_logout():
    script = (Path(__file__).parents[1] / "app" / "static" / "app.js").read_text(encoding="utf-8")
    sign_in = script.split("async function beginPkceLogin()", 1)[1].split("function sleep(", 1)[0]

    assert "force_reauth" not in sign_in
    assert "use_device_code:" in sign_in


def test_cli_login_status_without_active_login(monkeypatch):
    monkeypatch.setattr(main_module, "CLI_LOGIN_PROCESS", None)

    response = client.get("/api/auth/cli-login-status")

    assert response.json() == {"status": "not_started"}


def test_authentication_page_explains_cross_organization_login():
    response = client.get("/")

    assert response.status_code == 200
    assert 'id="authTenant"' in response.text
    assert 'id="deviceLoginCode"' in response.text
    assert "domain-joined PC is not required" in response.text
    assert response.text.index('<option value="browser">') < response.text.index('<option value="device">')
    assert "MFA steps requested by Microsoft" in response.text
    assert "no separate Azure CLI installation is required" in response.text


def test_cli_login_reuses_active_browser_flow(monkeypatch):
    class ActiveProcess:
        def poll(self):
            return None

    monkeypatch.setattr("app.main._resolve_az_command", lambda: "az.cmd")
    monkeypatch.setattr("app.main.CLI_LOGIN_PROCESS", ActiveProcess())

    response = client.post("/api/auth/cli-login")

    assert response.status_code == 200
    assert response.json()["status"] == "already_started"


def test_clear_auth_context_stops_active_login(monkeypatch):
    stopped = []
    monkeypatch.setattr("app.main._stop_cli_login_process", lambda: stopped.append(True))

    response = client.post("/api/auth/clear-context?clear_cli_session=false")

    assert response.status_code == 200
    assert stopped == [True]


def test_agent_report_returns_bounded_evidence(monkeypatch):
    async def fake_build_report(domain, payload, bearer_token):
        assert bearer_token == "agent-token"
        return ReportResult(
            domain=domain,
            tenant_id=payload.tenant_id,
            subscription_scope=payload.subscription_ids,
            generated_at_utc=datetime.now(UTC),
            execution_time_ms=12,
            data_quality_score=100.0,
            unresolved_fields_count=0,
            live_data_confirmed=True,
            source_audit=[],
            rows=[{"rank": 1}, {"rank": 2}],
            recommendations=[],
            warnings=[],
        )

    monkeypatch.setattr("app.main.build_report", fake_build_report)
    response = client.post(
        "/api/agent/reports/spend_visibility",
        headers={"Authorization": "Bearer agent-token"},
        json={
            "tenant_id": "tenant-a",
            "subscription_ids": ["sub-a"],
            "max_rows": 1,
        },
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["total_row_count"] == 2
    assert payload["returned_row_count"] == 1
    assert payload["result_truncated"] is True
    assert payload["rows"] == [{"rank": 1}]
    LAST_REPORT.clear()


def test_list_domains_includes_sample_inspired_live_reports():
    response = client.get("/api/domains")
    assert response.status_code == 200
    payload = response.json()
    assert len(payload) == 25
    assert {
        "tenant_spend_summary",
        "cost_summary",
        "reservation_savings",
        "quota_summary",
        "expiring_reservations",
        "storage_costs",
        "resource_group_costs",
        "subscription_cost_summary",
        "hybrid_use_benefit",
    } <= {item["key"] for item in payload}


def test_export_without_report_returns_404():
    response = client.get("/api/export/csv")
    assert response.status_code == 404


def test_repeated_export_reuses_cached_artifact(monkeypatch, tmp_path):
    report = ReportResult(
        domain=ReportDomain.optimization,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=100.0,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[],
        rows=[],
        recommendations=[],
        warnings=[],
    )
    output = tmp_path / "cached.csv"
    calls = 0

    def fake_exporter(_report):
        nonlocal calls
        calls += 1
        output.write_text("dataset\n", encoding="utf-8")
        return output

    LAST_REPORT.clear()
    LAST_REPORT.update({"report": report, "exports": {}})
    monkeypatch.setitem(main_module.EXPORT_MAP, "csv", fake_exporter)

    first = client.get("/api/export/csv")
    second = client.get("/api/export/csv")

    assert first.status_code == 200
    assert second.status_code == 200
    assert calls == 1


def test_executive_brief_uses_only_current_live_report_values():
    report = ReportResult(
        domain=ReportDomain.optimization,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=98.5,
        unresolved_fields_count=1,
        live_data_confirmed=True,
        source_audit=[],
        rows=[
            {"recommendationType": "Orphaned disk", "savingsEstimate": 125000},
            {"recommendationType": "Rightsize VM", "annualSavingsAmount": 75000},
        ],
        recommendations=[
            Recommendation(
                category="Commitment optimization",
                impact_estimate="$200,000",
                confidence="High",
                affected_scope="Selected subscriptions",
                suggested_action="Validate the reservation purchase recommendation.",
                evidence="Two live Advisor records.",
            )
        ],
        warnings=[],
    )

    brief = build_executive_brief(report)

    assert brief["title"] == "Azure FinOps Live Executive Report"
    assert ("Live potential savings", "$200,000.00") in brief["metrics"]
    assert ("Resources analyzed", "2") in brief["metrics"]
    assert ("Recommendations", "1") in brief["metrics"]
    assert brief["priorities"] == [
        (
            "Commitment optimization",
            "Validate the reservation purchase recommendation. Evidence: Two live Advisor records.",
        )
    ]
    rendered_brief = json.dumps(brief)
    assert "$1.2M" not in rendered_brief
    assert "$1.5M" not in rendered_brief
    assert "FY27" not in rendered_brief
    assert "90-Day" not in rendered_brief
    assert "All values are calculated from this report run" in brief["evidence_note"]


def test_pptx_tables_are_compact_wrapped_and_within_slide_bounds():
    report = ReportResult(
        domain=ReportDomain.optimization,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=98.5,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[
            SourceAudit(
                source="Azure Resource Graph",
                endpoint="https://management.azure.com/mock",
                duration_ms=10,
                row_count=1,
                executed_at_utc=datetime.now(UTC),
            )
        ],
        rows=[{
            "serviceScope": "All Compute Services",
            "subscriptionId": "sub-a",
            "resourceName": "Example virtual machine",
            "serviceName": "Compute",
            "resourceType": "Microsoft.Compute/virtualMachines",
            "resourceGroup": "rg-example",
            "region": "eastus",
            "cost": 42,
            "usageQuantity": 9,
            "estimatedTokenUnits": 0,
            "resourceId": "/subscriptions/sub-a/resourceGroups/rg-example/providers/Microsoft.Compute/virtualMachines/vm-example",
            "owner": "FinOps Team",
            "environment": "Production",
            "businessUnit": "Operations",
        }],
        recommendations=[
            Recommendation(
                category="Commitment optimization",
                impact_estimate="1 impacted resource",
                confidence="Medium",
                affected_scope="Selected subscriptions",
                suggested_action="Review the recommendation and assign an accountable owner.",
                evidence="Live query returned one result.",
            )
        ],
        warnings=[],
    )

    presentation = Presentation(export_pptx(report))
    slide_text = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    tables = [shape for slide in presentation.slides for shape in slide.shapes if shape.has_table]
    report_tables = [
        shape.table
        for slide in presentation.slides
        if any(
            shape.has_text_frame and shape.text.startswith("Report Results")
            for shape in slide.shapes
        )
        for shape in slide.shapes
        if shape.has_table
    ]

    assert presentation.slide_width > presentation.slide_height
    assert "Report Contents" in slide_text
    assert "Executive Overview" in slide_text
    assert "Recommendations" in slide_text
    assert "Report Results" in slide_text
    assert "Data Provenance" in slide_text
    assert any("1. Executive Overview" in text and "Slide 3" in text for text in slide_text)
    assert tables
    assert len(report_tables) > 1
    assert all(
        [table.cell(0, index).text for index in range(4)]
        == ["Row", "serviceScope", "subscriptionId", "resourceName"]
        for table in report_tables
    )
    report_headers = {
        table.cell(0, index).text
        for table in report_tables
        for index in range(len(table.columns))
    }
    assert {"owner", "environment", "businessUnit"} <= report_headers
    assert any(shape.height <= Inches(2) and len(shape.table.rows) == 2 for shape in tables)
    assert all(
        cell.text_frame.word_wrap
        and cell.text_frame.auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for shape in tables
        for row in shape.table.rows
        for cell in row.cells
    )
    assert all(
        shape.left >= 0
        and shape.top >= 0
        and shape.left + shape.width <= presentation.slide_width
        and shape.top + shape.height <= presentation.slide_height
        for slide in presentation.slides
        for shape in slide.shapes
    )
    priority_bodies = [
        shape.text_frame
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame and shape.text.startswith("Commitment optimization")
    ]
    assert priority_bodies
    assert priority_bodies[0].word_wrap
    assert priority_bodies[0].auto_size == MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE


def test_pdf_handles_unicode_and_contains_topic_outline():
    report = ReportResult(
        domain=ReportDomain.optimization,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=98.5,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[],
        rows=[{"resourceName": "Production VM – East", "cost": 42}],
        recommendations=[
            Recommendation(
                category="Cost optimization",
                impact_estimate="1 resource",
                confidence="High",
                affected_scope="Selected subscriptions",
                suggested_action="Review ‘idle’ usage and assign an owner.",
                evidence="Current-run evidence…",
            )
        ],
        warnings=["Partial access — one scope was unavailable."],
    )

    output = export_pdf(report)
    content = output.read_bytes()

    assert content.startswith(b"%PDF")
    assert len(content) > 1_000
    assert b"/Outlines" in content


def test_sample_inspired_domains_have_distinct_export_diagram_contracts():
    domains = [
        ReportDomain.tenant_spend_summary,
        ReportDomain.cost_summary,
        ReportDomain.budget_forecasting,
        ReportDomain.reservation_savings,
        ReportDomain.quota_summary,
        ReportDomain.expiring_reservations,
        ReportDomain.storage_costs,
        ReportDomain.resource_group_costs,
        ReportDomain.subscription_cost_summary,
        ReportDomain.savings_plans,
    ]
    signatures = []
    for domain in domains:
        report = ReportResult(
            domain=domain,
            tenant_id="tenant-a",
            subscription_scope=["sub-a"],
            generated_at_utc=datetime.now(UTC),
            execution_time_ms=10,
            data_quality_score=100.0,
            unresolved_fields_count=0,
            live_data_confirmed=False,
            source_audit=[],
            rows=[],
            recommendations=[],
            warnings=["No accessible live rows were returned."],
        )
        diagram = build_report_diagram(report)

        assert diagram is not None
        assert diagram["scope"]
        assert diagram["logic"]
        assert diagram["panels"]
        signatures.append(tuple(panel["title"] for panel in diagram["panels"]))

    assert len(set(signatures)) == len(domains)


def test_every_report_domain_has_export_diagrams():
    for domain in ReportDomain:
        report = ReportResult(
            domain=domain,
            tenant_id="tenant-a",
            subscription_scope=["sub-a"],
            generated_at_utc=datetime.now(UTC),
            execution_time_ms=10,
            data_quality_score=100.0,
            unresolved_fields_count=0,
            live_data_confirmed=False,
            source_audit=[],
            rows=[],
            recommendations=[],
            warnings=["No accessible live rows were returned."],
        )

        diagram = build_report_diagram(report)

        assert diagram is not None, domain.value
        assert diagram["scope"], domain.value
        assert diagram["logic"], domain.value
        assert diagram["panels"], domain.value


def test_cost_summary_exports_native_diagrams_and_explicit_empty_states():
    report = ReportResult(
        domain=ReportDomain.cost_summary,
        tenant_id="tenant-a",
        subscription_scope=["sub-a", "sub-b"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=100.0,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[],
        rows=[
            {
                "usageDate": "2026-07",
                "subscriptionId": "sub-a",
                "chargeType": "Usage",
                "meterCategory": "Compute",
                "cost": 120.0,
            },
            {
                "usageDate": "2026-08",
                "subscriptionId": "sub-b",
                "chargeType": "Purchase",
                "meterCategory": "Storage",
                "cost": 80.0,
            },
        ],
        recommendations=[],
        warnings=[],
    )

    presentation = Presentation(export_pptx(report))
    slide_text = [
        shape.text
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    charts = [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.has_chart
    ]
    pdf_content = export_pdf(report).read_bytes()

    assert "Live Report Diagram: Cost Summary" in slide_text
    assert {
        "Cost by subscription",
        "Subscription share",
        "Monthly charge-type mix",
        "Cost by meter category",
    } <= set(slide_text)
    assert len(charts) == 4
    assert pdf_content.count(b"/Type /Page") - pdf_content.count(b"/Type /Pages") >= 4

    empty_report = report.model_copy(
        update={
            "domain": ReportDomain.reservation_savings,
            "rows": [],
            "live_data_confirmed": False,
            "warnings": ["Reservation sources were inaccessible."],
        }
    )
    empty_presentation = Presentation(export_pptx(empty_report))
    empty_text = [
        shape.text
        for slide in empty_presentation.slides
        for shape in slide.shapes
        if shape.has_text_frame
    ]
    assert sum(text.startswith("No live values were returned") for text in empty_text) == 4


def test_export_with_report_returns_file():
    LAST_REPORT["report"] = ReportResult(
        domain=ReportDomain.ai_token_tracking,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=100.0,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[
            SourceAudit(
                source="Azure Cost Management",
                endpoint="https://management.azure.com/mock",
                duration_ms=10,
                row_count=1,
                executed_at_utc=datetime.now(UTC),
            )
        ],
        rows=[{"subscriptionId": "sub-a", "cost": 1.23, "resourceId": "/subscriptions/sub-a"}],
        enrichment_audit=[{"resolutionStatus": "resolved"}],
        ai_token_tracking=[{"serviceName": "Azure OpenAI", "usageQuantity": 100, "resourceId": "/subscriptions/sub-a"}],
        ai_token_summary=[{"serviceScope": "All AI Services", "records": 1, "totalUsageQuantity": 100, "averageUsageQuantity": 100, "totalEstimatedTokenUnits": 0, "averageEstimatedTokenUnits": 0, "totalCost": 1.23, "averageCost": 1.23}],
        recommendations=[
            Recommendation(
                category="AI workload efficiency",
                impact_estimate="1 impacted resources",
                confidence="Medium",
                affected_scope="Selected subscriptions",
                suggested_action="Review usage trends.",
                evidence="Live query returned one row",
            )
        ],
        warnings=[],
    )

    response = client.get("/api/export/csv")
    assert response.status_code == 200
    assert len(response.content) > 0
    assert int(response.headers["content-length"]) == len(response.content)

    response = client.get("/api/export/pptx")
    assert response.status_code == 200
    assert len(response.content) > 0
    assert int(response.headers["content-length"]) == len(response.content)
    assert response.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.presentationml.presentation")
    assert "report" in LAST_REPORT


def test_exports_do_not_flush_cached_report():
    LAST_REPORT["report"] = ReportResult(
        domain=ReportDomain.ai_token_tracking,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=100.0,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[
            SourceAudit(
                source="Azure Cost Management",
                endpoint="https://management.azure.com/mock",
                duration_ms=10,
                row_count=1,
                executed_at_utc=datetime.now(UTC),
            )
        ],
        rows=[{"subscriptionId": "sub-a", "cost": 1.23, "resourceId": "/subscriptions/sub-a"}],
        enrichment_audit=[{"resolutionStatus": "resolved"}],
        ai_token_tracking=[{"serviceName": "Azure OpenAI", "usageQuantity": 100, "resourceId": "/subscriptions/sub-a"}],
        ai_token_summary=[{"serviceScope": "All AI Services", "records": 1, "totalUsageQuantity": 100, "averageUsageQuantity": 100, "totalEstimatedTokenUnits": 0, "averageEstimatedTokenUnits": 0, "totalCost": 1.23, "averageCost": 1.23}],
        recommendations=[
            Recommendation(
                category="AI workload efficiency",
                impact_estimate="1 impacted resources",
                confidence="Medium",
                affected_scope="Selected subscriptions",
                suggested_action="Review usage trends.",
                evidence="Live query returned one row",
            )
        ],
        warnings=[],
    )

    csv_response = client.get("/api/export/csv")
    assert csv_response.status_code == 200
    assert "report" in LAST_REPORT

    pdf_response = client.get("/api/export/pdf")
    assert pdf_response.status_code == 200
    assert "report" in LAST_REPORT

    pptx_response = client.get("/api/export/pptx")
    assert pptx_response.status_code == 200
    assert "report" in LAST_REPORT


def test_clear_auth_context_clears_cached_report():
    LAST_REPORT["report"] = ReportResult(
        domain=ReportDomain.ai_token_tracking,
        tenant_id="tenant-a",
        subscription_scope=["sub-a"],
        generated_at_utc=datetime.now(UTC),
        execution_time_ms=10,
        data_quality_score=100.0,
        unresolved_fields_count=0,
        live_data_confirmed=True,
        source_audit=[
            SourceAudit(
                source="Azure Cost Management",
                endpoint="https://management.azure.com/mock",
                duration_ms=10,
                row_count=1,
                executed_at_utc=datetime.now(UTC),
            )
        ],
        rows=[{"subscriptionId": "sub-a", "cost": 1.23, "resourceId": "/subscriptions/sub-a"}],
        enrichment_audit=[{"resolutionStatus": "resolved"}],
        ai_token_tracking=[],
        ai_token_summary=[],
        recommendations=[
            Recommendation(
                category="AI workload efficiency",
                impact_estimate="1 impacted resources",
                confidence="Medium",
                affected_scope="Selected subscriptions",
                suggested_action="Review usage trends.",
                evidence="Live query returned one row",
            )
        ],
        warnings=[],
    )

    clear_response = client.post("/api/auth/clear-context?clear_cli_session=false")
    assert clear_response.status_code == 200
    assert clear_response.json().get("status") == "cleared"
    assert "report" not in LAST_REPORT

    export_response = client.get("/api/export/csv")
    assert export_response.status_code == 404


def test_governance_report_includes_policy_compliance(monkeypatch):
    async def fake_query_resource_graph(self, subscription_ids, query):
        return [], SourceAudit(
            source="Azure Resource Graph",
            endpoint="https://management.azure.com/mock",
            duration_ms=1,
            row_count=0,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_query_policy_compliance(self, subscription_ids):
        return [
            {
                "subscriptionId": "sub-a",
                "resourceId": "/subscriptions/sub-a/providers/Microsoft.PolicyInsights/policystates/latest",
                "policyAssignmentName": "enforce-tags",
                "policyDefinitionName": "require-tags",
                "complianceState": "NonCompliant",
                "timestamp": "2026-07-24T00:00:00Z",
            }
        ], SourceAudit(
            source="Azure Policy Compliance",
            endpoint="https://management.azure.com/mock-policy",
            duration_ms=2,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_resource_graph", fake_query_resource_graph)
    monkeypatch.setattr(AzureLiveClient, "query_policy_compliance", fake_query_policy_compliance)

    report = asyncio.run(
        build_report(
            ReportDomain.governance,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
                strict_live_mode=False,
            ),
            bearer_token="token",
        )
    )

    assert report.rows
    assert any(row.get("complianceState") == "NonCompliant" for row in report.rows)


def test_savings_potential_report_uses_advisor_data(monkeypatch):
    async def fake_list_advisor_recommendations(self, subscription_id):
        return [
            {
                "properties": {
                    "category": "Cost",
                    "impact": "Medium",
                    "shortDescription": {
                        "problem": "Right-size underutilized VM",
                        "solution": "Reduce SKU or shut down when idle",
                    },
                    "extendedProperties": {"annualSavingsAmount": 125.5},
                    "resourceMetadata": {"resourceId": "/subscriptions/sub-a/resourceGroups/rg/providers/Microsoft.Compute/virtualMachines/vm1"},
                }
            },
            {
                "properties": {
                    "category": "Cost",
                    "impact": "Low",
                    "shortDescription": {
                        "problem": "Purchase reservation for steady workload",
                        "solution": "Use a 1-year reservation",
                    },
                    "extendedProperties": {"annualSavingsAmount": 40.0},
                    "resourceMetadata": {"resourceId": "/subscriptions/sub-a/resourceGroups/rg/providers/Microsoft.Sql/servers/sql1"},
                }
            },
        ], SourceAudit(
            source="Azure Advisor",
            endpoint="https://management.azure.com/mock-advisor",
            duration_ms=2,
            row_count=2,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_query_ai_token_tracking(self, subscription_id, timeframe="MonthToDate"):
        return [], SourceAudit(
            source="Azure Cost Management (AI Token Tracking)",
            endpoint="https://management.azure.com/mock-ai",
            duration_ms=1,
            row_count=0,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "list_advisor_recommendations", fake_list_advisor_recommendations)
    monkeypatch.setattr(AzureLiveClient, "query_ai_token_tracking", fake_query_ai_token_tracking)

    report = asyncio.run(
        build_report(
            ReportDomain.savings_potential,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
                strict_live_mode=False,
            ),
            bearer_token="token",
        )
    )

    assert report.rows
    assert all("saving" in str(row.get("recommendationType", "")).lower() or "cost" in str(row.get("category", "")).lower() for row in report.rows)
    assert any(row.get("savingsEstimate") == 125.5 for row in report.rows)


def test_reservations_potential_report_uses_advisor_data(monkeypatch):
    async def fake_list_advisor_recommendations(self, subscription_id):
        return [
            {
                "properties": {
                    "category": "Cost",
                    "impact": "High",
                    "shortDescription": {
                        "problem": "Purchase reservation for steady SQL workload",
                        "solution": "Use a 1-year reservation",
                    },
                    "extendedProperties": {"annualSavingsAmount": 99.0},
                    "resourceMetadata": {"resourceId": "/subscriptions/sub-a/resourceGroups/rg/providers/Microsoft.Sql/servers/sql1"},
                }
            },
            {
                "properties": {
                    "category": "Cost",
                    "impact": "Low",
                    "shortDescription": {
                        "problem": "Reduce app service scale out",
                        "solution": "Right-size the instance",
                    },
                    "extendedProperties": {"annualSavingsAmount": 12.0},
                    "resourceMetadata": {"resourceId": "/subscriptions/sub-a/resourceGroups/rg/providers/Microsoft.Web/sites/app1"},
                }
            },
        ], SourceAudit(
            source="Azure Advisor",
            endpoint="https://management.azure.com/mock-advisor",
            duration_ms=2,
            row_count=2,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_query_ai_token_tracking(self, subscription_id, timeframe="MonthToDate"):
        return [], SourceAudit(
            source="Azure Cost Management (AI Token Tracking)",
            endpoint="https://management.azure.com/mock-ai",
            duration_ms=1,
            row_count=0,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "list_advisor_recommendations", fake_list_advisor_recommendations)
    monkeypatch.setattr(AzureLiveClient, "query_ai_token_tracking", fake_query_ai_token_tracking)

    report = asyncio.run(
        build_report(
            ReportDomain.reservations_potential,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
                strict_live_mode=False,
            ),
            bearer_token="token",
        )
    )

    assert report.rows
    assert any("reservation" in str(row.get("recommendationType", "")).lower() for row in report.rows)
    assert any(row.get("savingsEstimate") == 99.0 for row in report.rows)


def test_savings_plan_recommendation_query_uses_azure_benefit_api(monkeypatch):
    captured = {}

    class FakeResponse:
        def raise_for_status(self):
            return None

        def json(self):
            return {"value": [{"kind": "SavingsPlan"}]}

    class FakeAsyncClient:
        def __init__(self, timeout):
            captured["timeout"] = timeout

        async def __aenter__(self):
            return self

        async def __aexit__(self, exc_type, exc, traceback):
            return None

        async def get(self, url, headers, params):
            captured["url"] = url
            captured["params"] = params
            captured["headers"] = headers
            return FakeResponse()

    monkeypatch.setattr("app.services.azure_live.httpx.AsyncClient", FakeAsyncClient)

    rows, audit = asyncio.run(
        AzureLiveClient("test-token").query_savings_plan_recommendations("sub-a", "P1Y")
    )

    assert rows == [{"kind": "SavingsPlan"}]
    assert captured["url"] == (
        "https://management.azure.com/subscriptions/sub-a/providers/"
        "Microsoft.CostManagement/benefitRecommendations"
    )
    assert captured["params"] == {
        "api-version": "2026-06-01",
        "$filter": (
            "properties/scope eq 'Single' AND "
            "properties/lookBackPeriod eq 'Last30Days' AND "
            "properties/term eq 'P1Y'"
        ),
    }
    assert audit.source == "Azure Cost Management Savings Plan Recommendations"


def test_savings_plans_report_uses_cost_management_benefit_recommendations(monkeypatch):
    queried_terms = []

    async def fake_query_savings_plan_recommendations(self, subscription_id, term):
        queried_terms.append((subscription_id, term))
        recommendations = [
            {
                "kind": "SavingsPlan",
                "name": f"{term}-lower",
                "properties": {
                    "armSkuName": "Compute_Savings_Plan",
                    "scope": "Single",
                    "term": term,
                    "lookBackPeriod": "Last30Days",
                    "firstConsumptionDate": "2026-09-01T00:00:00Z",
                    "lastConsumptionDate": "2026-09-30T00:00:00Z",
                    "totalHours": 720,
                    "commitmentGranularity": "Hourly",
                    "costWithoutBenefit": 1000.0,
                    "currencyCode": "USD",
                    "recommendationDetails": {
                        "commitmentAmount": 1.25,
                        "savingsAmount": 25.0,
                        "savingsPercentage": 2.5,
                        "coveragePercentage": 50.0,
                        "averageUtilizationPercentage": 95.0,
                        "wastageCost": 0.1,
                    },
                },
            },
            {
                "kind": "SavingsPlan",
                "name": f"{term}-highest",
                "properties": {
                    "armSkuName": "Compute_Savings_Plan",
                    "scope": "Single",
                    "term": term,
                    "lookBackPeriod": "Last30Days",
                    "totalHours": 720,
                    "commitmentGranularity": "Hourly",
                    "costWithoutBenefit": 1000.0,
                    "currencyCode": "USD",
                    "recommendationDetails": {
                        "commitmentAmount": 1.5,
                        "savingsAmount": 45.0 if term == "P1Y" else 70.0,
                        "savingsPercentage": 4.5,
                        "coveragePercentage": 60.0,
                        "averageUtilizationPercentage": 92.0,
                        "wastageCost": 0.2,
                    },
                },
            },
            {
                "kind": "Reservation",
                "name": f"{term}-reservation",
                "properties": {
                    "recommendationDetails": {"savingsAmount": 500.0},
                },
            },
        ]
        return recommendations, SourceAudit(
            source="Azure Cost Management Savings Plan Recommendations",
            endpoint="https://management.azure.com/mock-benefit-recommendations",
            duration_ms=2,
            row_count=len(recommendations),
            executed_at_utc=datetime.now(UTC),
        )

    async def unexpected_cost_query(self, subscription_id, timeframe="MonthToDate"):
        raise AssertionError("Savings Plans must not use the current-cost baseline query")

    monkeypatch.setattr(
        AzureLiveClient,
        "query_savings_plan_recommendations",
        fake_query_savings_plan_recommendations,
    )
    monkeypatch.setattr(AzureLiveClient, "query_cost_management", unexpected_cost_query)

    report = asyncio.run(
        build_report(
            ReportDomain.savings_plans,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
                filters={
                    "start_date": "2026-09-01",
                    "end_date": "2026-09-30",
                },
                strict_live_mode=False,
            ),
            bearer_token="token",
        )
    )

    assert queried_terms == [("sub-a", "P1Y"), ("sub-a", "P3Y")]
    assert len(report.rows) == 4
    assert all(row["recordType"] == "savings_plan_recommendation" for row in report.rows)
    top_options = [row for row in report.rows if row["isTopSavingsOption"]]
    assert {(row["term"], row["savingsAmount"]) for row in top_options} == {
        ("P1Y", 45.0),
        ("P3Y", 70.0),
    }
    assert report.data_quality_score == 100.0
    assert len(report.recommendations) == 2
    assert any("trailing 30-day usage lookback" in warning for warning in report.warnings)
    assert all(audit.source == "Azure Cost Management Savings Plan Recommendations" for audit in report.source_audit)


def test_ai_token_tracking_uses_filter_dates_and_dict_rows(monkeypatch):
    captured: dict[str, object] = {}

    async def fake_query_ai_token_tracking(self, subscription_id, timeframe="MonthToDate", start_date=None, end_date=None):
        captured["subscription_id"] = subscription_id
        captured["start_date"] = start_date
        captured["end_date"] = end_date
        return [
            {
                "UsageQuantity": 250,
                "totalCost": 3.5,
                "ResourceId": "/subscriptions/sub-a/resourceGroups/rg/providers/Microsoft.CognitiveServices/accounts/aoai1",
                "ServiceName": "Azure OpenAI",
                "Meter": "InputToken",
            }
        ], SourceAudit(
            source="Azure Cost Management (AI Token Tracking)",
            endpoint="https://management.azure.com/mock-ai",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_query_resource_graph(self, subscription_ids, query):
        return [], SourceAudit(
            source="Azure Resource Graph",
            endpoint="https://management.azure.com/mock-rg",
            duration_ms=1,
            row_count=0,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_ai_token_tracking", fake_query_ai_token_tracking)
    monkeypatch.setattr(AzureLiveClient, "query_resource_graph", fake_query_resource_graph)

    report = asyncio.run(
        build_report(
            ReportDomain.ai_token_tracking,
            __import__("app.models", fromlist=["ReportRequest", "FilterSet"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
                strict_live_mode=False,
                filters=__import__("app.models", fromlist=["FilterSet"]).FilterSet(
                    start_date="2026-01-01",
                    end_date="2026-07-24",
                ),
            ),
            bearer_token="token",
        )
    )

    assert captured["subscription_id"] == "sub-a"
    assert captured["start_date"] == "2026-01-01"
    assert captured["end_date"] == "2026-07-24"
    assert report.ai_token_tracking
    assert report.ai_token_tracking[0]["serviceName"] == "Azure OpenAI"
    assert report.ai_token_tracking[0]["usageQuantity"] == 250


def test_tenant_spend_summary_aggregates_live_cost_rows_and_honors_dates(monkeypatch):
    captured = {}

    async def fake_query_cost_details(self, subscription_id, start_date=None, end_date=None):
        captured[subscription_id] = (start_date, end_date)
        return [
            {
                "UsageDate": "2026-08-01",
                "PreTaxCost": 10,
                "UsageQuantity": 2,
                "MeterCategory": "Virtual Machines",
                "MeterSubCategory": "Dv5 Series",
                "Meter": "D4 v5",
                "UnitOfMeasure": "1 Hour",
                "ServiceName": "Virtual Machines",
                "ChargeType": "Usage",
            },
            {
                "UsageDate": "2026-08-01",
                "PreTaxCost": 5,
                "UsageQuantity": 3,
                "MeterCategory": "Virtual Machines",
                "MeterSubCategory": "Dv5 Series",
                "Meter": "D4 v5",
                "UnitOfMeasure": "1 Hour",
                "ServiceName": "Virtual Machines",
                "ChargeType": "Usage",
            },
        ], SourceAudit(
            source="Azure Cost Management",
            endpoint="https://management.azure.com/live-cost",
            duration_ms=1,
            row_count=2,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_cost_details", fake_query_cost_details)
    report = asyncio.run(
        build_report(
            ReportDomain.tenant_spend_summary,
            __import__("app.models", fromlist=["ReportRequest", "FilterSet"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a", "sub-b"],
                filters=__import__("app.models", fromlist=["FilterSet"]).FilterSet(
                    start_date="2026-08-01",
                    end_date="2026-08-31",
                ),
            ),
            bearer_token="token",
        )
    )

    assert captured == {
        "sub-a": ("2026-08-01", "2026-08-31"),
        "sub-b": ("2026-08-01", "2026-08-31"),
    }
    assert report.rows == [
        {
            "usageDate": "2026-08-01",
            "meterCategory": "Virtual Machines",
            "meterSubCategory": "Dv5 Series",
            "meterName": "D4 v5",
            "unitOfMeasure": "1 Hour",
            "serviceName": "Virtual Machines",
            "chargeType": "Usage",
            "subscriptionId": "sub-a",
            "cost": 15.0,
            "usageQuantity": 5.0,
            "recordCount": 2,
        },
        {
            "usageDate": "2026-08-01",
            "meterCategory": "Virtual Machines",
            "meterSubCategory": "Dv5 Series",
            "meterName": "D4 v5",
            "unitOfMeasure": "1 Hour",
            "serviceName": "Virtual Machines",
            "chargeType": "Usage",
            "subscriptionId": "sub-b",
            "cost": 15.0,
            "usageQuantity": 5.0,
            "recordCount": 2,
        },
    ]
    assert report.live_data_confirmed is True


def test_tenant_spend_summary_keeps_accessible_rows_when_other_subscriptions_are_denied(monkeypatch):
    async def fake_query_cost_details(self, subscription_id, start_date=None, end_date=None):
        if subscription_id == "sub-denied":
            raise RuntimeError("403 Cost Management Reader access denied")
        return [
            {
                "UsageDate": "2026-08-01",
                "PreTaxCost": 25,
                "UsageQuantity": 1,
                "MeterCategory": "Virtual Machines",
                "MeterSubCategory": "Dv5 Series",
                "Meter": "D4 v5",
                "UnitOfMeasure": "1 Hour",
                "ServiceName": "Virtual Machines",
                "ChargeType": "Usage",
            }
        ], SourceAudit(
            source="Azure Cost Management",
            endpoint="https://management.azure.com/live-cost",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_cost_details", fake_query_cost_details)
    report = asyncio.run(
        build_report(
            ReportDomain.tenant_spend_summary,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-accessible", "sub-denied"],
            ),
            bearer_token="token",
        )
    )

    assert len(report.rows) == 1
    assert report.rows[0]["subscriptionId"] == "sub-accessible"
    assert report.rows[0]["cost"] == 25.0
    assert report.live_data_confirmed is True
    assert any("sub-denied" in warning for warning in report.warnings)


def test_scope_options_returns_success_when_some_selected_subscriptions_are_denied(monkeypatch):
    async def fake_query_resource_graph(self, subscription_ids, query):
        raise RuntimeError("403 Resource Graph access denied")

    async def fake_list_advisor_recommendations(self, subscription_id):
        if subscription_id == "sub-denied":
            raise RuntimeError("403 Advisor access denied")
        return [{"properties": {"impact": "High"}}], SourceAudit(
            source="Azure Advisor",
            endpoint="https://management.azure.com/live-advisor",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_resource_graph", fake_query_resource_graph)
    monkeypatch.setattr(AzureLiveClient, "list_advisor_recommendations", fake_list_advisor_recommendations)

    response = client.post(
        "/api/scope-options",
        headers={"Authorization": "Bearer token"},
        json={"subscription_ids": ["sub-accessible", "sub-denied"]},
    )

    assert response.status_code == 200
    payload = response.json()
    assert payload["severity"] == ["High"]
    assert len(payload["warnings"]) == 2


def test_quota_summary_uses_live_compute_usage_and_calculates_utilization(monkeypatch):
    async def fake_query_resource_graph(self, subscription_ids, query):
        return [{"subscriptionId": "sub-a", "location": "eastus"}], SourceAudit(
            source="Azure Resource Graph",
            endpoint="https://management.azure.com/live-resource-graph",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_list_compute_usages(self, subscription_id, location):
        return [{
            "subscriptionId": subscription_id,
            "region": location,
            "quotaName": "Total Regional vCPUs",
            "currentValue": 80,
            "limit": 100,
            "unit": "Count",
        }], SourceAudit(
            source="Azure Compute Usage",
            endpoint="https://management.azure.com/live-usage",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_resource_graph", fake_query_resource_graph)
    monkeypatch.setattr(AzureLiveClient, "list_compute_usages", fake_list_compute_usages)
    report = asyncio.run(
        build_report(
            ReportDomain.quota_summary,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
            ),
            bearer_token="token",
        )
    )

    assert report.rows[0]["utilizationPercent"] == 80.0
    assert report.rows[0]["remaining"] == 20.0
    assert {audit.source for audit in report.source_audit} == {
        "Azure Resource Graph",
        "Azure Compute Usage",
    }


def test_expiring_reservations_filters_to_selected_scope_and_calculates_days(monkeypatch):
    async def fake_query_reservation_costs(self, subscription_id, start_date=None, end_date=None):
        return [
            {
                "recordType": "reservation_cost",
                "subscriptionId": subscription_id,
                "reservationId": "selected",
                "cost": 125.0,
            },
            {
                "recordType": "reservation_cost",
                "subscriptionId": subscription_id,
                "reservationId": "selected-prefix-only",
                "cost": 999.0,
            },
        ], SourceAudit(
            source="Azure Cost Management (Reservations)",
            endpoint="https://management.azure.com/live-reservation-costs",
            duration_ms=1,
            row_count=2,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_list_reservations(self):
        return [
            {
                "reservationId": "selected",
                "expiryDate": "2026-09-30T00:00:00Z",
                "appliedScopes": ["/subscriptions/sub-a"],
            },
            {
                "reservationId": "other",
                "expiryDate": "2026-09-30T00:00:00Z",
                "appliedScopes": ["/subscriptions/sub-b"],
            },
        ], SourceAudit(
            source="Azure Reservations",
            endpoint="https://management.azure.com/live-reservations",
            duration_ms=1,
            row_count=2,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_reservation_costs", fake_query_reservation_costs)
    monkeypatch.setattr(AzureLiveClient, "list_reservations", fake_list_reservations)
    report = asyncio.run(
        build_report(
            ReportDomain.expiring_reservations,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
            ),
            bearer_token="token",
        )
    )

    assert [row["reservationId"] for row in report.rows] == ["selected"]
    assert report.rows[0]["recordType"] == "reservation_inventory"
    assert report.rows[0]["incurredCost"] == 125.0
    assert isinstance(report.rows[0]["remainingDays"], int)


def test_reservation_savings_keeps_cost_and_advisor_evidence_distinct(monkeypatch):
    async def fake_query_reservation_costs(self, subscription_id, start_date=None, end_date=None):
        return [{
            "recordType": "reservation_cost",
            "subscriptionId": subscription_id,
            "reservationId": "ri-1",
            "reservationName": "Production compute",
            "cost": 75.0,
        }], SourceAudit(
            source="Azure Cost Management (Reservations)",
            endpoint="https://management.azure.com/live-reservation-costs",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    async def fake_list_advisor_recommendations(self, subscription_id):
        return [{
            "properties": {
                "category": "Cost",
                "impact": "High",
                "shortDescription": {
                    "problem": "Purchase reservation for steady compute",
                    "solution": "Use a 1-year reservation",
                },
                "extendedProperties": {"annualSavingsAmount": 240.0},
            }
        }], SourceAudit(
            source="Azure Advisor",
            endpoint="https://management.azure.com/live-advisor",
            duration_ms=1,
            row_count=1,
            executed_at_utc=datetime.now(UTC),
        )

    monkeypatch.setattr(AzureLiveClient, "query_reservation_costs", fake_query_reservation_costs)
    monkeypatch.setattr(AzureLiveClient, "list_advisor_recommendations", fake_list_advisor_recommendations)
    report = asyncio.run(
        build_report(
            ReportDomain.reservation_savings,
            __import__("app.models", fromlist=["ReportRequest"]).ReportRequest(
                tenant_id="tenant-a",
                subscription_ids=["sub-a"],
            ),
            bearer_token="token",
        )
    )

    assert {row["recordType"] for row in report.rows} == {
        "advisor_recommendation",
        "reservation_cost",
    }
    assert sum(row.get("cost", 0) for row in report.rows) == 75.0
    assert sum(float(row.get("savingsEstimate", 0)) for row in report.rows) == 240.0
