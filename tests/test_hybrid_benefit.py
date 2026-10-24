# Copyright (c) 2026 jerry wolff

import asyncio
import csv
import json
import shutil
import sqlite3
import subprocess
from datetime import UTC, datetime
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient
from pptx import Presentation

from app.main import LAST_REPORT, app
from app.models import FilterSet, ReportDomain, ReportRequest, SourceAudit
from app.services.azure_live import AzureLiveClient
from app.services.diagrams import build_report_diagram
from app.services.exporters import export_csv, export_pdf, export_pptx, export_sqlite
from app.services.reporting import _normalize_hybrid_benefit_row, build_report


def inventory_row(name="test-vm", **overrides):
    """Isolated synthetic inventory; never used by production collection."""
    return {
        "id": f"/subscriptions/test-sub/resourceGroups/test-rg/providers/Microsoft.Compute/virtualMachines/{name}",
        "subscriptionId": "test-sub",
        "name": name,
        "type": "microsoft.compute/virtualmachines",
        "resourceGroup": "test-rg",
        "location": "eastus",
        "osType": "Windows",
        "licenseType": "Windows_Server",
        "tags": {"environment": "test", "costCenter": "example"},
        **overrides,
    }


def audit(count):
    return SourceAudit(
        source="Azure Resource Graph",
        endpoint="https://management.azure.com/providers/Microsoft.ResourceGraph/resources",
        duration_ms=1,
        row_count=count,
        executed_at_utc=datetime.now(UTC),
    )


@pytest.fixture
def live_inventory(monkeypatch):
    async def validate(self, subscription_id, tenant_id):
        assert tenant_id == "test-tenant"
        assert subscription_id in {"test-sub", "denied-sub"}

    async def collect(self, subscription_id):
        if subscription_id == "denied-sub":
            response = httpx.Response(
                403, request=httpx.Request("GET", "https://management.azure.com"),
            )
            raise httpx.HTTPStatusError("sensitive-error-must-not-appear", request=response.request, response=response)
        rows = [
            inventory_row("enabled"),
            inventory_row("candidate", licenseType="None"),
            inventory_row("unknown", licenseType=None),
        ]
        return rows, audit(len(rows))

    monkeypatch.setattr(AzureLiveClient, "validate_subscription_tenant", validate)
    monkeypatch.setattr(AzureLiveClient, "query_hybrid_benefit", collect)


@pytest.mark.parametrize(
    "overrides,status",
    [
        ({}, "Enabled"),
        ({"type": "microsoft.compute/virtualmachinescalesets"}, "Enabled"),
        ({"licenseType": "None"}, "Not enabled"),
        ({"licenseType": None}, "Unknown"),
        ({"licenseType": "unexpected"}, "Unknown"),
        ({"osType": None}, "Unknown"),
        ({"licenseType": "Windows_Client"}, "Not applicable"),
        ({"osType": "Linux"}, "Not applicable"),
        ({"type": "microsoft.sql/servers/databases", "licenseType": "BasePrice"}, "Enabled"),
        ({"type": "microsoft.sql/managedinstances", "licenseType": "BasePrice"}, "Enabled"),
        ({"type": "microsoft.sql/servers/elasticpools", "licenseType": "LicenseIncluded"}, "Not enabled"),
        ({"type": "microsoft.sql/servers/databases", "licenseType": "LicenseIncluded"}, "Not enabled"),
        ({"type": "microsoft.sql/servers/databases", "skuTier": "Standard"}, "Not applicable"),
        ({"type": "microsoft.sql/servers/databases", "skuName": "GP_S_Gen5"}, "Not applicable"),
        ({"type": "microsoft.sql/servers/databases", "name": "server/master"}, "Not applicable"),
        ({"type": "microsoft.sql/servers/databases", "elasticPoolId": "test-pool"}, "Inherited (pool)"),
        ({"type": "microsoft.sqlvirtualmachine/sqlvirtualmachines", "sqlServerLicenseType": "AHUB"}, "Enabled"),
        ({"type": "microsoft.sqlvirtualmachine/sqlvirtualmachines", "sqlServerLicenseType": "PAYG"}, "Not enabled"),
        ({"type": "microsoft.sqlvirtualmachine/sqlvirtualmachines", "sqlServerLicenseType": "DR"}, "Not applicable"),
        ({"type": "microsoft.sqlvirtualmachine/sqlvirtualmachines"}, "Unknown"),
    ],
)
def test_license_status_is_observed_not_assumed(overrides, status):
    row = _normalize_hybrid_benefit_row(inventory_row(**overrides))
    assert row["benefitStatus"] == status
    assert row["reviewCandidate"] is (status == "Not enabled")
    assert row["statusReason"]
    assert not {"savingsEstimate", "ownedLicenses", "cost"} & row.keys()


def test_hybrid_report_preserves_partial_coverage_and_provenance(live_inventory):
    report = asyncio.run(build_report(
        ReportDomain.hybrid_use_benefit,
        ReportRequest(
            tenant_id="test-tenant", subscription_ids=["test-sub", "denied-sub"],
            filters=FilterSet(start_date="2026-10-01"),
        ),
        "synthetic-test-token",
    ))
    assert report.live_data_confirmed
    assert len(report.rows) == 3
    assert len(report.source_audit) == 1
    assert report.source_audit[0].row_count == 3
    assert all(row["sourceService"] == "Azure Resource Graph" and row["collectedAtUtc"] for row in report.rows)
    assert report.unresolved_fields_count == 1
    assert report.data_quality_score < 100
    assert any("HTTP 403" in warning and "Reader" in warning for warning in report.warnings)
    assert any("Date filters do not apply" in warning for warning in report.warnings)
    assert "sensitive-error-must-not-appear" not in report.model_dump_json()
    assert [row["resourceName"] for row in report.rows] == ["candidate", "enabled", "unknown"]
    assert len(report.recommendations) == 2
    assert all("Not quantified" in item.impact_estimate for item in report.recommendations)


@pytest.mark.parametrize("filters,expected", [
    (FilterSet(actionable_only=True), ["candidate"]),
    (FilterSet(region="EASTUS", resource_group="TEST-RG", tag_key="environment", tag_value="test"), ["candidate", "enabled", "unknown"]),
    (FilterSet(region="westus"), []),
    (FilterSet(tag_key="missing"), []),
    (FilterSet(environment="prod"), []),
    (FilterSet(cost_center="example"), ["candidate", "enabled", "unknown"]),
    (FilterSet(resource_type="microsoft.sql/managedinstances"), []),
])
def test_hybrid_filters(live_inventory, filters, expected):
    report = asyncio.run(build_report(
        ReportDomain.hybrid_use_benefit,
        ReportRequest(tenant_id="test-tenant", subscription_ids=["test-sub"], filters=filters),
        "synthetic-test-token",
    ))
    assert [row["resourceName"] for row in report.rows] == expected


def test_empty_denied_and_invalid_scope_are_distinct(live_inventory, monkeypatch):
    request = ReportRequest(tenant_id="test-tenant", subscription_ids=["denied-sub"])
    denied = asyncio.run(build_report(ReportDomain.hybrid_use_benefit, request, "test-token"))
    assert not denied.live_data_confirmed
    assert not denied.rows
    assert any("HTTP 403" in warning for warning in denied.warnings)

    async def empty(self, subscription_id):
        return [], audit(0)

    monkeypatch.setattr(AzureLiveClient, "query_hybrid_benefit", empty)
    request.subscription_ids = ["test-sub"]
    empty_report = asyncio.run(build_report(ReportDomain.hybrid_use_benefit, request, "test-token"))
    assert empty_report.live_data_confirmed
    assert not empty_report.rows
    assert not any("HTTP" in warning for warning in empty_report.warnings)

    async def invalid(self, subscription_id, tenant_id):
        raise ValueError("sensitive-tenant-details")

    async def should_not_collect(self, subscription_id):
        pytest.fail("Inventory must not run before tenant validation")

    monkeypatch.setattr(AzureLiveClient, "validate_subscription_tenant", invalid)
    monkeypatch.setattr(AzureLiveClient, "query_hybrid_benefit", should_not_collect)
    invalid_report = asyncio.run(build_report(ReportDomain.hybrid_use_benefit, request, "test-token"))
    assert not invalid_report.live_data_confirmed
    assert any("scope or response validation failed" in warning for warning in invalid_report.warnings)
    assert "sensitive-tenant-details" not in invalid_report.model_dump_json()


def test_resource_graph_pagination_and_hybrid_query(monkeypatch):
    calls = []

    async def post(self, url, **kwargs):
        calls.append(json.loads(json.dumps(kwargs["json"])))
        page = {"data": [inventory_row(str(len(calls)))]}
        if len(calls) == 1:
            page["$skipToken"] = "page-two"
            page["resultTruncated"] = "true"
        return httpx.Response(200, json=page, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx.AsyncClient, "post", post)
    rows, source = asyncio.run(AzureLiveClient("test-token").query_hybrid_benefit("test-sub"))
    assert len(rows) == source.row_count == 2
    assert calls[0]["subscriptions"] == ["test-sub"]
    assert calls[1]["options"]["$skipToken"] == "page-two"
    assert calls[0]["options"]["resultFormat"] == "objectArray"
    for field in ("licenseType", "sqlServerLicenseType", "elasticPoolId", "virtualMachineProfile", "order by id"):
        assert field in calls[0]["query"]


@pytest.mark.parametrize("page", [
    {"data": [], "resultTruncated": "true"},
    {"data": [], "$skipToken": "repeated-token"},
])
def test_truncated_inventory_never_looks_complete(monkeypatch, page):
    async def post(self, url, **kwargs):
        return httpx.Response(200, json=page, request=httpx.Request("POST", url))

    monkeypatch.setattr(httpx.AsyncClient, "post", post)
    with pytest.raises(RuntimeError):
        asyncio.run(AzureLiveClient("test-token").query_hybrid_benefit("test-sub"))


@pytest.mark.parametrize("actual_tenant", ["test-tenant", "other-tenant", None])
def test_scope_validation_uses_delegated_arm_response(monkeypatch, actual_tenant):
    async def get(self, url, **kwargs):
        assert "/subscriptions/test-sub?" in url
        return httpx.Response(
            200, json={"tenantId": actual_tenant}, request=httpx.Request("GET", url),
        )

    monkeypatch.setattr(httpx.AsyncClient, "get", get)
    operation = AzureLiveClient("test-token").validate_subscription_tenant("test-sub", "test-tenant")
    if actual_tenant == "test-tenant":
        asyncio.run(operation)
    else:
        with pytest.raises(ValueError, match="does not match"):
            asyncio.run(operation)


def test_hybrid_api_agent_exports_and_badge(live_inventory, monkeypatch, tmp_path):
    monkeypatch.setattr("app.services.exporters.EXPORT_DIR", tmp_path)
    client = TestClient(app)
    body = {"tenant_id": "test-tenant", "subscription_ids": ["test-sub"]}
    headers = {"Authorization": "Bearer synthetic-test-token"}
    LAST_REPORT.clear()
    try:
        assert client.post("/api/reports/hybrid_use_benefit", json=body).status_code == 401
        response = client.post("/api/reports/hybrid_use_benefit", json=body, headers=headers)
        assert response.status_code == 200
        report = LAST_REPORT["report"]
        assert response.json()["rows"] == report.rows
        diagram = build_report_diagram(report)
        assert len(diagram["panels"]) == 3
        assert diagram["panels"][0]["valueFormat"] == "number"
        assert sum(entry["value"] for entry in diagram["panels"][0]["entries"]) == 3

        csv_path = export_csv(report)
        with csv_path.open(encoding="utf-8", newline="") as source:
            csv_rows = list(csv.DictReader(source))
        context = json.loads(next(row["reportContext"] for row in csv_rows if row["dataset"] == "report_context"))
        assert context["warnings"] == report.warnings
        assert context["source_audit"][0]["row_count"] == 3
        evidence_rows = [row for row in csv_rows if row["dataset"] == "report_rows"]
        assert {row["benefitStatus"] for row in evidence_rows} == {"Enabled", "Not enabled", "Unknown"}
        assert all(row["statusReason"] and row["collectedAtUtc"] for row in evidence_rows)
        sqlite_path = export_sqlite(report)
        with sqlite3.connect(sqlite_path) as connection:
            stored = [json.loads(row[0]) for row in connection.execute("SELECT row_json FROM report_rows")]
            assert stored == report.rows
            assert connection.execute("SELECT COUNT(*) FROM source_audit").fetchone()[0] == 1
            sqlite_context = json.loads(connection.execute("SELECT context_json FROM report_context").fetchone()[0])
            assert sqlite_context == context
        assert export_pdf(report).read_bytes().startswith(b"%PDF")
        presentation = Presentation(export_pptx(report))
        text = " ".join(shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame)
        text += " ".join(
            cell.text
            for slide in presentation.slides for shape in slide.shapes if shape.has_table
            for row in shape.table.rows for cell in row.cells
        )
        assert "Hybrid Benefit configuration" in text
        assert "Not quantified" in text

        agent = client.post("/api/agent/reports/hybrid_use_benefit", json={**body, "max_rows": 1}, headers=headers)
        assert agent.status_code == 200
        assert agent.json()["total_row_count"] == 3
        assert agent.json()["returned_row_count"] == 1
        assert agent.json()["result_truncated"]
        assert agent.json()["source_audit"]
        assert agent.json()["warnings"]
        html = client.get("/").text
        assert 'class="copilot-icon"' in html
        assert "does not invoke Copilot or an MCP server" in html
        assert client.get("/static/copilot-badge.svg").status_code == 200
    finally:
        LAST_REPORT.clear()


def test_core_sources_include_copyright():
    root = Path(__file__).parents[1]
    sources = [
        *root.joinpath("app").rglob("*.py"),
        *root.joinpath("app", "static").glob("*"),
        root / "app" / "templates" / "index.html",
        root / "function_app.py",
        root / "desktop_launcher.py",
        *root.glob("*.ps1"),
        root / "m365-agent" / "build-package.ps1",
        root / "Dockerfile",
        root / "FinOpsCCoE.spec",
        root / "installer" / "FinOpsCCoE.iss",
    ]
    for source in sources:
        assert "Copyright (c) 2026 jerry wolff" in source.read_text(encoding="utf-8"), source.name


@pytest.mark.skipif(shutil.which("node") is None, reason="Node.js is required for browser KPI contract validation")
def test_hybrid_browser_kpis_use_configuration_counts():
    root = Path(__file__).parents[1]
    script = r"""
const assert = require("node:assert/strict");
const fs = require("node:fs");
const vm = require("node:vm");
const source = fs.readFileSync("app\\static\\app.js", "utf8");
const start = source.indexOf("function buildVisualReport(data)");
const end = source.indexOf("function renderResults(data)", start);
const kpis = [];
const context = {
  document: { createElement: () => ({ append() {}, appendChild() {} }) },
  visualReportDomains: new Set(["hybrid_use_benefit"]),
  humanizeFieldName: String,
  reportDiagramSpecs: { hybrid_use_benefit: {scope:"test", logic:"test"} },
  escapeHtml: String,
  numericValue: Number,
  formatCompactNumber: String,
  addKpi: (_, label, value) => kpis.push([label, value]),
  renderDomainDiagrams() {},
};
vm.createContext(context);
vm.runInContext(source.slice(start, end), context);
for (const rows of [
  [{benefitStatus:"Enabled"}, {benefitStatus:"Not enabled", reviewCandidate:true}, {benefitStatus:"Unknown"}],
  [],
]) {
  kpis.length = 0;
  context.buildVisualReport({domain:"hybrid_use_benefit", rows, source_audit:[],
    subscription_scope:[], generated_at_utc:"2026-10-24T00:00:00Z", data_quality_score:100});
  const expected = rows.length ? "1" : "0";
  assert.deepEqual(kpis.slice(0, 3), [
    ["Benefit enabled", expected], ["Review candidates", expected], ["Unknown configuration", expected],
  ]);
  assert.ok(!kpis.some(([label]) => /cost|savings/i.test(label)));
}
"""
    result = subprocess.run(
        [shutil.which("node"), "-"], input=script, cwd=root,
        capture_output=True, text=True, check=False,
    )
    assert result.returncode == 0, result.stdout + result.stderr
