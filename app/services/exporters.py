from __future__ import annotations

import csv
import json
import os
import re
import sqlite3
from pathlib import Path
import sys

from fpdf import FPDF
from pptx import Presentation
from pptx.chart.data import ChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.util import Inches, Pt

from app.models import ReportResult
from app.services.diagrams import build_report_diagram


if getattr(sys, "frozen", False):
    EXPORT_DIR = Path(os.environ.get("LOCALAPPDATA", Path.home())) / "FinOpsCCoE" / "exports"
else:
    EXPORT_DIR = Path("exports")


def _pdf_safe_text(text: object, max_chunk: int = 40) -> str:
    raw = str(text).translate(
        str.maketrans(
            {
                "\u2013": "-",
                "\u2014": "-",
                "\u2018": "'",
                "\u2019": "'",
                "\u201c": '"',
                "\u201d": '"',
                "\u2022": "-",
                "\u2026": "...",
                "\u00a0": " ",
            }
        )
    ).encode("latin-1", "replace").decode("latin-1")
    # Insert soft breaks into long non-whitespace runs (resource IDs, URLs, GUID chains)
    # so FPDF can always render within cell width.
    return re.sub(
        rf"\S{{{max_chunk + 1},}}",
        lambda m: " ".join(
            m.group(0)[i:i + max_chunk] for i in range(0, len(m.group(0)), max_chunk)
        ),
        raw,
    )


def _pdf_table_text(value: object, max_chunk: int = 36) -> str:
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return _pdf_safe_text(cleaned, max_chunk=max_chunk)


def _pptx_safe_text(value: object, max_chunk: int = 28) -> str:
    cleaned = re.sub(r"\s+", " ", str(value)).strip()
    return re.sub(
        rf"\S{{{max_chunk + 1},}}",
        lambda match: " ".join(
            match.group(0)[index:index + max_chunk]
            for index in range(0, len(match.group(0)), max_chunk)
        ),
        cleaned,
    )


def ensure_export_dir() -> None:
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)


def _safe_number(value: object) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _format_money(value: float) -> str:
    return f"${value:,.2f}"


def build_executive_brief(report: ReportResult) -> dict[str, object]:
    total_cost = sum(_safe_number(row.get("cost")) for row in report.rows)
    total_savings = sum(
        _safe_number(row.get("savingsEstimate") or row.get("annualSavingsAmount"))
        for row in report.rows
    )
    priorities = [
        (
            recommendation.category,
            f"{recommendation.suggested_action} Evidence: {recommendation.evidence}",
        )
        for recommendation in report.recommendations[:4]
    ]
    if not priorities:
        priorities = [
            (
                "Validate report coverage",
                report.warnings[0] if report.warnings else "The live query returned no recommendation records.",
            )
        ]

    actions = [
        (f"Priority {index}", recommendation.suggested_action)
        for index, recommendation in enumerate(report.recommendations[:6], start=1)
    ]
    if not actions:
        actions = [("Priority 1", "Validate permissions, selected scope, and report filters before acting.")]

    decisions = [
        f"Assign an accountable owner for {recommendation.category.lower()} in {recommendation.affected_scope}."
        for recommendation in report.recommendations[:6]
    ]
    if not decisions:
        decisions = ["Confirm the selected tenant and subscription coverage before making a FinOps decision."]

    return {
        "title": "Azure FinOps Live Executive Report",
        "subtitle": report.domain.value.replace("_", " ").title(),
        "headline": (
            f"Review {len(report.rows):,} current-run result(s) from {len(report.source_audit):,} live Azure source call(s)."
        ),
        "metrics": [
            ("Live potential savings", _format_money(total_savings) if total_savings > 0 else "Not available in this report"),
            ("Cost represented", _format_money(total_cost) if total_cost > 0 else "Not available in this report"),
            ("Resources analyzed", f"{len(report.rows):,}"),
            ("Recommendations", f"{len(report.recommendations):,}"),
            ("Data quality", f"{report.data_quality_score:.1f}%"),
            ("Partial-access warnings", f"{len(report.warnings):,}"),
        ],
        "priorities": priorities,
        "actions": actions,
        "executive_decisions": decisions,
        "evidence_note": (
            f"Based on {len(report.rows)} live result(s), {len(report.recommendations)} recommendation(s), "
            f"and {len(report.source_audit)} source call(s) collected at "
            f"{report.generated_at_utc.strftime('%Y-%m-%d %H:%M UTC')}. All values are calculated from this report run."
        ),
    }


def export_csv(report: ReportResult) -> Path:
    ensure_export_dir()
    output = EXPORT_DIR / f"{report.domain.value}_{report.generated_at_utc.strftime('%Y%m%d_%H%M%S')}.csv"
    base_rows = [
        {"dataset": "report_rows", **row}
        for row in report.rows
    ]
    ai_rows = [
        {"dataset": "ai_token_tracking", **row}
        for row in report.ai_token_tracking
    ]
    ai_summary_rows = [
        {"dataset": "ai_token_summary", **row}
        for row in report.ai_token_summary
    ]
    rows = base_rows + ai_summary_rows + ai_rows
    headers = sorted(set().union(*(r.keys() for r in rows))) if rows else ["dataset"]
    with output.open("w", encoding="utf-8", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=headers)
        writer.writeheader()
        for row in rows:
            normalized = {
                key: json.dumps(value) if isinstance(value, (dict, list)) else value
                for key, value in row.items()
            }
            writer.writerow(normalized)
    return output


def export_sqlite(report: ReportResult) -> Path:
    ensure_export_dir()
    output = EXPORT_DIR / f"{report.domain.value}_{report.generated_at_utc.strftime('%Y%m%d_%H%M%S')}.sqlite"
    conn = sqlite3.connect(output)
    cur = conn.cursor()

    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS report_runs (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            tenant_id TEXT,
            generated_at_utc TEXT,
            execution_time_ms INTEGER,
            data_quality_score REAL,
            unresolved_fields_count INTEGER
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS report_rows (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            row_json TEXT
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS recommendations (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            category TEXT,
            impact_estimate TEXT,
            confidence TEXT,
            affected_scope TEXT,
            suggested_action TEXT,
            evidence TEXT
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS source_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            source TEXT,
            endpoint TEXT,
            duration_ms INTEGER,
            row_count INTEGER,
            executed_at_utc TEXT
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS enrichment_audit (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            row_json TEXT
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_token_tracking (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            row_json TEXT
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS ai_token_summary (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            domain TEXT,
            row_json TEXT
        )
        """
    )

    cur.execute(
        """
        INSERT INTO report_runs (
            domain, tenant_id, generated_at_utc, execution_time_ms,
            data_quality_score, unresolved_fields_count
        ) VALUES (?, ?, ?, ?, ?, ?)
        """,
        (
            report.domain.value,
            report.tenant_id,
            report.generated_at_utc.isoformat(),
            report.execution_time_ms,
            report.data_quality_score,
            report.unresolved_fields_count,
        ),
    )

    for row in report.rows:
        cur.execute(
            "INSERT INTO report_rows (domain, row_json) VALUES (?, ?)",
            (report.domain.value, json.dumps(row)),
        )

    for rec in report.recommendations:
        cur.execute(
            """
            INSERT INTO recommendations (
                domain, category, impact_estimate, confidence,
                affected_scope, suggested_action, evidence
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                report.domain.value,
                rec.category,
                rec.impact_estimate,
                rec.confidence,
                rec.affected_scope,
                rec.suggested_action,
                rec.evidence,
            ),
        )

    for audit in report.source_audit:
        cur.execute(
            """
            INSERT INTO source_audit (
                domain, source, endpoint, duration_ms, row_count, executed_at_utc
            ) VALUES (?, ?, ?, ?, ?, ?)
            """,
            (
                report.domain.value,
                audit.source,
                audit.endpoint,
                audit.duration_ms,
                audit.row_count,
                audit.executed_at_utc.isoformat(),
            ),
        )

    for row in report.enrichment_audit:
        cur.execute(
            "INSERT INTO enrichment_audit (domain, row_json) VALUES (?, ?)",
            (report.domain.value, json.dumps(row)),
        )

    for row in report.ai_token_tracking:
        cur.execute(
            "INSERT INTO ai_token_tracking (domain, row_json) VALUES (?, ?)",
            (report.domain.value, json.dumps(row)),
        )

    for row in report.ai_token_summary:
        cur.execute(
            "INSERT INTO ai_token_summary (domain, row_json) VALUES (?, ?)",
            (report.domain.value, json.dumps(row)),
        )

    conn.commit()
    conn.close()
    return output


DIAGRAM_PDF_COLORS = [
    (22, 138, 173),
    (15, 118, 110),
    (234, 88, 12),
    (202, 138, 4),
    (37, 99, 235),
    (190, 24, 93),
]


def _format_diagram_value(value: float, value_format: str) -> str:
    if value_format == "currency":
        return _format_money(value)
    if value_format == "percent":
        return f"{value:,.1f}%"
    if value_format == "days":
        return f"{value:,.0f} days"
    return f"{value:,.2f}" if value % 1 else f"{value:,.0f}"


def _draw_pdf_panel(
    pdf: FPDF,
    panel: dict[str, object],
    x: float,
    y: float,
    width: float,
    height: float,
) -> None:
    pdf.set_fill_color(245, 247, 250)
    pdf.set_draw_color(203, 213, 225)
    pdf.rect(x, y, width, height, style="DF")
    pdf.set_text_color(20, 31, 56)
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_xy(x + 3, y + 2.5)
    pdf.cell(width - 6, 4, _pdf_safe_text(panel["title"])[:70])
    pdf.set_text_color(71, 85, 105)
    pdf.set_font("Helvetica", "", 6.5)
    pdf.set_xy(x + 3, y + 7)
    pdf.cell(width - 6, 3.5, _pdf_safe_text(panel["subtitle"])[:105])

    kind = str(panel["kind"])
    value_format = str(panel.get("valueFormat", "number"))
    entries = list(panel.get("entries", []))[:10]
    categories = list(panel.get("categories", []))
    series = [str(value) for value in panel.get("series", [])]
    has_values = bool(categories and series) if kind == "share" else bool(entries)
    if kind == "donut":
        entries = [entry for entry in entries if float(entry["value"]) > 0]
        has_values = bool(entries)
    if kind == "share":
        has_values = has_values and any(
            any(float(value) != 0 for value in category["values"])
            for category in categories
        )

    if not has_values:
        pdf.set_font("Helvetica", "", 8)
        pdf.set_text_color(71, 85, 105)
        pdf.set_xy(x + 8, y + 24)
        pdf.multi_cell(
            width - 16,
            5,
            "No live values were returned for this diagram. Review Operational Notes for scope or access details.",
            align="C",
        )
        return

    if kind == "bars":
        shown = entries[:6]
        max_value = max(abs(float(entry["value"])) for entry in shown) or 1
        row_height = (height - 16) / max(len(shown), 1)
        label_width = min(43, width * 0.34)
        value_width = min(27, width * 0.22)
        track_width = width - label_width - value_width - 9
        for index, entry in enumerate(shown):
            row_y = y + 13 + (index * row_height)
            value = float(entry["value"])
            pdf.set_font("Helvetica", "", 6.5)
            pdf.set_text_color(51, 65, 85)
            pdf.set_xy(x + 3, row_y)
            pdf.cell(label_width, row_height, _pdf_table_text(entry["label"], 18)[:30])
            pdf.set_fill_color(226, 232, 240)
            pdf.rect(x + 4 + label_width, row_y + 1.4, track_width, max(2.5, row_height - 3), style="F")
            pdf.set_fill_color(22, 138, 173)
            pdf.rect(
                x + 4 + label_width,
                row_y + 1.4,
                track_width * (abs(value) / max_value),
                max(2.5, row_height - 3),
                style="F",
            )
            pdf.set_font("Helvetica", "B", 6.5)
            pdf.set_xy(x + 6 + label_width + track_width, row_y)
            pdf.cell(value_width, row_height, _format_diagram_value(value, value_format), align="R")
        return

    if kind == "columns":
        shown = entries[:8]
        max_value = max(abs(float(entry["value"])) for entry in shown) or 1
        chart_x = x + 4
        chart_y = y + 15
        chart_width = width - 8
        chart_height = height - 27
        slot_width = chart_width / max(len(shown), 1)
        pdf.set_draw_color(148, 163, 184)
        pdf.line(chart_x, chart_y + chart_height, chart_x + chart_width, chart_y + chart_height)
        for index, entry in enumerate(shown):
            value = float(entry["value"])
            bar_height = chart_height * (abs(value) / max_value)
            bar_width = max(2, slot_width * 0.62)
            bar_x = chart_x + (index * slot_width) + ((slot_width - bar_width) / 2)
            pdf.set_fill_color(22, 138, 173)
            pdf.rect(bar_x, chart_y + chart_height - bar_height, bar_width, bar_height, style="F")
            pdf.set_font("Helvetica", "", 5.5)
            pdf.set_text_color(71, 85, 105)
            pdf.set_xy(chart_x + (index * slot_width), chart_y + chart_height + 1)
            pdf.cell(slot_width, 3, _pdf_table_text(entry["label"], 10)[:10], align="C")
            pdf.set_xy(chart_x + (index * slot_width), max(chart_y, chart_y + chart_height - bar_height - 3))
            pdf.cell(slot_width, 3, _format_diagram_value(value, value_format)[:15], align="C")
        return

    if kind == "donut":
        shown = entries[:6]
        total = sum(float(entry["value"]) for entry in shown)
        diameter = min(height - 20, width * 0.28)
        donut_x = x + 5
        donut_y = y + 15
        angle = 0.0
        for index, entry in enumerate(shown):
            sweep = (float(entry["value"]) / total) * 360
            pdf.set_fill_color(*DIAGRAM_PDF_COLORS[index % len(DIAGRAM_PDF_COLORS)])
            pdf.solid_arc(donut_x, donut_y, diameter, angle, angle + sweep, style="F")
            angle += sweep
        hole_size = diameter * 0.45
        pdf.set_fill_color(245, 247, 250)
        pdf.ellipse(
            donut_x + (diameter / 2) - (hole_size / 2),
            donut_y + (diameter / 2) - (hole_size / 2),
            hole_size,
            hole_size,
            style="F",
        )
        legend_x = donut_x + diameter + 5
        legend_width = width - (legend_x - x) - 4
        for index, entry in enumerate(shown):
            legend_y = y + 14 + (index * 7)
            pdf.set_fill_color(*DIAGRAM_PDF_COLORS[index % len(DIAGRAM_PDF_COLORS)])
            pdf.rect(legend_x, legend_y + 1, 3, 3, style="F")
            pdf.set_font("Helvetica", "", 6.5)
            pdf.set_text_color(51, 65, 85)
            pdf.set_xy(legend_x + 5, legend_y)
            share = (float(entry["value"]) / total) * 100
            pdf.cell(
                legend_width - 5,
                5,
                f"{_pdf_table_text(entry['label'], 18)[:24]}  {share:.1f}%",
            )
        return

    if kind == "share":
        shown_categories = categories[:5]
        shown_series = series[:6]
        label_width = min(32, width * 0.28)
        track_width = width - label_width - 8
        row_height = (height - 25) / max(len(shown_categories), 1)
        for category_index, category in enumerate(shown_categories):
            row_y = y + 14 + (category_index * row_height)
            values = [float(value) for value in category["values"][:len(shown_series)]]
            total = sum(values)
            pdf.set_font("Helvetica", "", 6.5)
            pdf.set_text_color(51, 65, 85)
            pdf.set_xy(x + 3, row_y)
            pdf.cell(label_width, row_height, _pdf_table_text(category["label"], 12)[:18])
            segment_x = x + 4 + label_width
            for series_index, value in enumerate(values):
                segment_width = track_width * (value / total) if total else 0
                pdf.set_fill_color(*DIAGRAM_PDF_COLORS[series_index % len(DIAGRAM_PDF_COLORS)])
                pdf.rect(segment_x, row_y + 1.5, segment_width, max(3, row_height - 3), style="F")
                segment_x += segment_width
        legend_x = x + 4
        legend_y = y + height - 7
        for series_index, series_name in enumerate(shown_series):
            item_width = (width - 8) / max(len(shown_series), 1)
            pdf.set_fill_color(*DIAGRAM_PDF_COLORS[series_index % len(DIAGRAM_PDF_COLORS)])
            pdf.rect(legend_x + (series_index * item_width), legend_y + 1, 3, 3, style="F")
            pdf.set_font("Helvetica", "", 5.5)
            pdf.set_xy(legend_x + (series_index * item_width) + 4, legend_y)
            pdf.cell(item_width - 4, 5, _pdf_table_text(series_name, 10)[:14])


def _add_pdf_diagram_page(
    pdf: FPDF,
    report: ReportResult,
    diagram: dict[str, object],
    palette: dict[str, tuple[int, int, int]],
) -> None:
    pdf.add_page()
    pdf.start_section("Report Diagram", level=0)
    page_width = pdf.w - pdf.l_margin - pdf.r_margin
    pdf.set_fill_color(*palette["navy"])
    pdf.rect(0, 0, pdf.w, 20, style="F")
    pdf.set_text_color(*palette["white"])
    pdf.set_font("Helvetica", "B", 15)
    pdf.set_xy(pdf.l_margin, 6)
    pdf.cell(page_width * 0.65, 7, _pdf_safe_text(f"Live Report Diagram: {diagram['title']}"))
    pdf.set_font("Helvetica", "", 7)
    pdf.set_xy(pdf.l_margin + (page_width * 0.65), 7)
    pdf.cell(
        page_width * 0.35,
        5,
        f"Collected {report.generated_at_utc.strftime('%Y-%m-%d %H:%M UTC')}",
        align="R",
    )

    band_y = 24
    band_height = 17
    pdf.set_fill_color(240, 253, 250)
    pdf.set_draw_color(203, 213, 225)
    pdf.rect(pdf.l_margin, band_y, page_width, band_height, style="DF")
    definition_width = (page_width - 4) / 2
    pdf.set_text_color(*palette["slate"])
    pdf.set_font("Helvetica", "B", 7)
    pdf.set_xy(pdf.l_margin + 3, band_y + 2)
    pdf.cell(definition_width - 6, 4, "SCOPE")
    pdf.set_font("Helvetica", "", 7)
    pdf.set_xy(pdf.l_margin + 3, band_y + 6)
    pdf.multi_cell(definition_width - 6, 3.5, _pdf_table_text(diagram["scope"], 42))
    logic_x = pdf.l_margin + definition_width + 4
    pdf.set_font("Helvetica", "B", 7)
    pdf.set_xy(logic_x, band_y + 2)
    pdf.cell(definition_width - 6, 4, "LOGIC")
    pdf.set_font("Helvetica", "", 7)
    pdf.set_xy(logic_x, band_y + 6)
    pdf.multi_cell(definition_width - 6, 3.5, _pdf_table_text(diagram["logic"], 42))

    panel_gap = 4
    panel_width = (page_width - panel_gap) / 2
    panel_height = 69
    positions = [
        (pdf.l_margin, 46),
        (pdf.l_margin + panel_width + panel_gap, 46),
        (pdf.l_margin, 119),
        (pdf.l_margin + panel_width + panel_gap, 119),
    ]
    for panel, (x, y) in zip(list(diagram["panels"])[:4], positions, strict=False):
        _draw_pdf_panel(pdf, panel, x, y, panel_width, panel_height)

    pdf.set_text_color(*palette["slate"])


def export_pdf(report: ReportResult) -> Path:
    ensure_export_dir()
    output = EXPORT_DIR / f"{report.domain.value}_{report.generated_at_utc.strftime('%Y%m%d_%H%M%S')}.pdf"
    brief = build_executive_brief(report)

    pdf = FPDF(orientation="L")
    pdf.set_margins(10, 10, 10)
    pdf.add_page()
    pdf.set_auto_page_break(auto=True, margin=15)

    palette = {
        "navy": (20, 31, 56),
        "teal": (15, 118, 110),
        "sky": (226, 232, 240),
        "slate": (51, 65, 85),
        "white": (255, 255, 255),
    }

    page_width = pdf.w - pdf.l_margin - pdf.r_margin

    def ensure_space(required_height: float) -> None:
        if pdf.get_y() + required_height > pdf.h - 15:
            pdf.add_page()

    def section_title(title: str, *, register: bool = True) -> None:
        ensure_space(12)
        if register:
            pdf.start_section(_pdf_safe_text(title), level=0)
        pdf.set_fill_color(*palette["teal"])
        pdf.set_text_color(*palette["white"])
        pdf.set_font("Helvetica", "B", 11)
        pdf.set_x(pdf.l_margin)
        pdf.cell(page_width, 8, _pdf_safe_text(title), new_x="LMARGIN", new_y="NEXT", fill=True)
        pdf.set_text_color(*palette["slate"])

    def topic_header(title: str, context: str) -> None:
        ensure_space(28)
        if pdf.get_y() > pdf.t_margin + 1:
            pdf.ln(10)
        section_title(title)
        pdf.set_font("Helvetica", "I", 8)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 4.5, _pdf_table_text(context, 48))
        pdf.ln(1)

    def render_toc(toc_pdf: FPDF, outline: list[object]) -> None:
        toc_pdf.set_fill_color(*palette["navy"])
        toc_pdf.rect(0, 0, toc_pdf.w, 24, style="F")
        toc_pdf.set_text_color(*palette["white"])
        toc_pdf.set_font("Helvetica", "B", 18)
        toc_pdf.set_xy(toc_pdf.l_margin, 8)
        toc_pdf.cell(page_width, 8, "Table of Contents")
        toc_pdf.set_text_color(*palette["slate"])
        toc_pdf.set_y(32)
        toc_pdf.set_font("Helvetica", "", 10)
        for item in outline:
            if item.level != 0:
                continue
            link = toc_pdf.add_link(page=item.page_number)
            toc_pdf.set_x(toc_pdf.l_margin)
            toc_pdf.cell(page_width - 18, 7, _pdf_safe_text(item.name), link=link)
            toc_pdf.cell(18, 7, str(item.page_number), align="R", link=link, new_x="LMARGIN", new_y="NEXT")

    def draw_table(title: str, columns: list[tuple[str, float]], rows: list[dict[str, object]]) -> None:
        section_title(title, register=False)
        ensure_space(8)
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_fill_color(*palette["sky"])
        for label, width in columns:
            pdf.cell(width, 6, _pdf_table_text(label, 24)[:36], border=1, fill=True)
        pdf.ln()

        pdf.set_font("Helvetica", "", 7)
        for row in rows:
            ensure_space(6)
            for key, width in columns:
                text = _pdf_table_text(row.get(key, ""), 32)[:140]
                pdf.cell(width, 6, text, border=1)
            pdf.ln()

    # Header band
    pdf.set_fill_color(*palette["navy"])
    pdf.rect(0, 0, pdf.w, 32, style="F")
    pdf.set_text_color(*palette["white"])
    pdf.set_font("Helvetica", "B", 18)
    pdf.set_xy(pdf.l_margin, 10)
    pdf.cell(page_width, 8, _pdf_safe_text(brief["title"]), new_x="LMARGIN", new_y="NEXT")
    pdf.set_font("Helvetica", "", 10)
    pdf.cell(page_width, 6, _pdf_safe_text(brief["subtitle"]), new_x="LMARGIN", new_y="NEXT")

    pdf.set_text_color(*palette["slate"])
    pdf.set_y(39)
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(page_width, 6, _pdf_table_text(brief["headline"], 48))
    pdf.ln(2)
    pdf.set_font("Helvetica", "", 8)
    pdf.multi_cell(
        page_width,
        5,
        _pdf_table_text(
            f"Tenant: {report.tenant_id} | Subscriptions: {len(report.subscription_scope)} | "
            f"Generated: {report.generated_at_utc.strftime('%Y-%m-%d %H:%M UTC')}",
            48,
        ),
    )

    pdf.add_page()
    pdf.insert_toc_placeholder(render_toc, pages=1)

    diagram = build_report_diagram(report)
    if diagram:
        _add_pdf_diagram_page(pdf, report, diagram, palette)
        pdf.add_page()

    # Executive summary
    topic_header("Executive Summary", "Current-run financial, coverage, and data-quality indicators for executive review.")
    pdf.set_font("Helvetica", "", 10)
    pdf.multi_cell(page_width, 6, _pdf_table_text(brief["headline"], 40))
    for label, value in brief["metrics"]:
        ensure_space(6)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 6, _pdf_table_text(f"{label}: {value}", 40))

    topic_header("Executive Priorities", "Priority outcomes framed by the report context and supported only by current-run evidence.")
    pdf.set_font("Helvetica", "", 9)
    for title, detail in brief["priorities"]:
        ensure_space(8)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 5.5, _pdf_table_text(f"{title}: {detail}", 40))

    topic_header("90-Day Action Plan", "Sequenced actions for validating, assigning, and realizing the opportunities in this report.")
    pdf.set_font("Helvetica", "", 9)
    for period, action in brief["actions"]:
        ensure_space(8)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 5.5, _pdf_table_text(f"{period}: {action}", 40))

    topic_header("Executive Decisions", "Decisions and ownership required to move report findings into accountable execution.")
    pdf.set_font("Helvetica", "", 9)
    for decision in brief["executive_decisions"]:
        ensure_space(7)
        pdf.set_x(pdf.l_margin)
        pdf.multi_cell(page_width, 5.5, _pdf_table_text(f"- {decision}", 40))
    ensure_space(8)
    pdf.set_x(pdf.l_margin)
    pdf.multi_cell(page_width, 5.5, _pdf_table_text(brief["evidence_note"], 40))

    def row_summary(row: dict[str, object]) -> str:
        if row.get("recommendationType"):
            parts = [str(row.get("recommendationType", ""))]
            if row.get("category"):
                parts.append(str(row.get("category")))
            if row.get("impact"):
                parts.append(f"impact {row.get('impact')}")
            if row.get("savingsEstimate") not in (None, ""):
                parts.append(f"estimated savings {_format_money(_safe_number(row.get('savingsEstimate')))}")
            return " | ".join(parts)

        parts = []
        for field in ("resourceName", "serviceName", "resourceGroup", "region"):
            value = row.get(field)
            if value not in (None, ""):
                parts.append(f"{field}: {_pdf_table_text(value, 18)}")
        if row.get("cost") not in (None, ""):
            parts.append(f"cost {_format_money(_safe_number(row.get('cost')))}")
        if row.get("usageQuantity") not in (None, ""):
            parts.append(f"usage {_pdf_table_text(row.get('usageQuantity'), 18)}")
        if row.get("estimatedTokenUnits") not in (None, ""):
            parts.append(f"tokens {_pdf_table_text(row.get('estimatedTokenUnits'), 18)}")
        if row.get("savingsEstimate") not in (None, ""):
            parts.append(f"estimated savings {_format_money(_safe_number(row.get('savingsEstimate')))}")
        return " | ".join(parts) or "Result item"

    # Recommendations table
    recommendation_rows = [
        {
            "Category": rec.category,
            "Impact": rec.impact_estimate,
            "Confidence": rec.confidence,
            "Action": rec.suggested_action,
        }
        for rec in report.recommendations
    ]
    if recommendation_rows:
        topic_header("Recommendations", "Prioritized findings, impact, confidence, and accountable remediation actions.")
        draw_table(
            "Top Recommendations",
            [("Category", 38), ("Impact", 36), ("Confidence", 24), ("Action", page_width - 98)],
            recommendation_rows[:8],
        )

    # AI summary table
    if report.ai_token_summary:
        topic_header("AI Usage", "Aggregated Azure AI consumption, token estimates, and current-run cost evidence.")
        ai_summary_rows = [
            {
                "Scope": row.get("serviceScope", ""),
                "Total": row.get("totalUsageQuantity", ""),
                "Avg": row.get("averageUsageQuantity", ""),
                "Tokens": row.get("totalEstimatedTokenUnits", ""),
                "AvgTok": row.get("averageEstimatedTokenUnits", ""),
                "Cost": row.get("totalCost", ""),
            }
            for row in report.ai_token_summary
        ]
        draw_table(
            "AI Usage Summary",
            [("Scope", 44), ("Total", 24), ("Avg", 24), ("Tokens", 24), ("AvgTok", 24), ("Cost", page_width - 140)],
            ai_summary_rows[:10],
        )

    # Report results table
    if report.rows:
        topic_header("Report Results", "Representative live results for the selected report topic and subscription scope.")
        section_title("Results Highlights", register=False)
        pdf.set_font("Helvetica", "", 8)
        highlight_lines = [row_summary(row) for row in report.rows[:8]]
        for line in highlight_lines:
            ensure_space(6)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(page_width, 5.5, _pdf_table_text(f"- {line}", 40))

    if report.source_audit or report.enrichment_audit:
        topic_header("Data Provenance", "Source execution and enrichment evidence supporting this report.")
        if report.source_audit:
            source_rows = [
                {
                    "Source": item.source,
                    "Rows": item.row_count,
                    "Duration": f"{item.duration_ms} ms",
                    "Executed": item.executed_at_utc.strftime("%Y-%m-%d %H:%M UTC"),
                }
                for item in report.source_audit
            ]
            draw_table(
                "Source Audit",
                [("Source", 70), ("Rows", 24), ("Duration", 32), ("Executed", page_width - 126)],
                source_rows[:10],
            )
        if report.enrichment_audit:
            pdf.set_font("Helvetica", "", 8)
            ensure_space(7)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(
                page_width,
                5,
                f"Enrichment audit records: {len(report.enrichment_audit):,} | "
                f"Unresolved key fields: {report.unresolved_fields_count:,}",
            )

    # Footer notes
    if report.warnings:
        topic_header("Operational Notes", "Warnings, partial-access conditions, and limitations affecting interpretation.")
        section_title("Warnings", register=False)
        pdf.set_font("Helvetica", "", 8)
        for warning in report.warnings[:6]:
            ensure_space(6)
            pdf.set_x(pdf.l_margin)
            pdf.multi_cell(page_width, 6, _pdf_table_text(f"- {warning}", 40))

    pdf.output(str(output))
    return output


def export_pptx(report: ReportResult) -> Path:
    ensure_export_dir()
    output = EXPORT_DIR / f"{report.domain.value}_{report.generated_at_utc.strftime('%Y%m%d_%H%M%S')}.pptx"
    brief = build_executive_brief(report)

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    palette = {
        "navy": RGBColor(20, 31, 56),
        "teal": RGBColor(15, 118, 110),
        "slate": RGBColor(51, 65, 85),
        "pale": RGBColor(245, 247, 250),
        "light": RGBColor(226, 232, 240),
        "white": RGBColor(255, 255, 255),
    }
    section_slides: dict[str, int] = {}

    def set_text(
        text_frame,
        text: str,
        *,
        size: int = 18,
        bold: bool = False,
        color: RGBColor | None = None,
        align: PP_ALIGN = PP_ALIGN.LEFT,
    ) -> None:
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        paragraph = text_frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.alignment = align
        run = paragraph.runs[0]
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color or palette["slate"]

    def add_title_slide() -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        banner = slide.shapes.add_shape(
            1, Inches(0), Inches(0), presentation.slide_width, Inches(1.3)
        )
        banner.fill.solid()
        banner.fill.fore_color.rgb = palette["navy"]
        banner.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.3), Inches(0.5))
        set_text(title_box.text_frame, str(brief["title"]), size=26, bold=True, color=palette["white"])

        domain_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.75), Inches(12.3), Inches(0.35))
        set_text(
            domain_box.text_frame,
            str(brief["subtitle"]),
            size=12,
            color=palette["white"],
        )

        summary_box = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.0), Inches(1.4))
        summary_lines = [
            f"Tenant: {report.tenant_id}",
            f"Subscriptions: {len(report.subscription_scope)}",
            f"Generated: {report.generated_at_utc.isoformat()}",
            f"Execution: {report.execution_time_ms} ms | Data Quality: {report.data_quality_score}%",
        ]
        text_frame = summary_box.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for idx, line in enumerate(summary_lines):
            paragraph = text_frame.paragraphs[0] if idx == 0 else text_frame.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            if paragraph.runs:
                run = paragraph.runs[0]
                run.font.size = Pt(14 if idx == 0 else 12)
                run.font.color.rgb = palette["slate"]

    def add_summary_slide() -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        header = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.5))
        set_text(header.text_frame, "Executive Summary", size=22, bold=True, color=palette["navy"])

        bullets = [str(brief["headline"])] + [
            f"{label}: {value}" for label, value in brief["metrics"]
        ] + [str(brief["evidence_note"])]

        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(4.8))
        tf = body.text_frame
        tf.clear()
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for idx, line in enumerate(bullets):
            paragraph = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            if paragraph.runs:
                run = paragraph.runs[0]
                run.font.size = Pt(16 if idx == 0 else 13)
                run.font.color.rgb = palette["slate"]

    def add_chart_panel(
        slide,
        panel: dict[str, object],
        left: float,
        top: float,
        width: float,
        height: float,
    ) -> None:
        background = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(left),
            Inches(top),
            Inches(width),
            Inches(height),
        )
        background.fill.solid()
        background.fill.fore_color.rgb = palette["pale"]
        background.line.color.rgb = palette["light"]

        title_box = slide.shapes.add_textbox(
            Inches(left + 0.18), Inches(top + 0.1), Inches(width - 0.36), Inches(0.25)
        )
        set_text(title_box.text_frame, str(panel["title"]), size=11, bold=True, color=palette["navy"])
        subtitle_box = slide.shapes.add_textbox(
            Inches(left + 0.18), Inches(top + 0.37), Inches(width - 0.36), Inches(0.2)
        )
        set_text(subtitle_box.text_frame, str(panel["subtitle"]), size=7, color=palette["slate"])

        kind = str(panel["kind"])
        chart_data = ChartData()
        has_values = False
        if kind == "share":
            series = [str(value) for value in panel.get("series", [])]
            categories = list(panel.get("categories", []))
            chart_data.categories = [str(category["label"]) for category in categories]
            for series_index, series_name in enumerate(series):
                values = [float(category["values"][series_index]) for category in categories]
                chart_data.add_series(series_name, values)
                has_values = has_values or any(value != 0 for value in values)
            chart_type = XL_CHART_TYPE.BAR_STACKED_100
        else:
            entries = list(panel.get("entries", []))[:10]
            if kind == "donut":
                entries = [entry for entry in entries if float(entry["value"]) > 0]
            chart_data.categories = [str(entry["label"]) for entry in entries]
            values = [float(entry["value"]) for entry in entries]
            if entries:
                chart_data.add_series(str(panel["title"]), values)
            has_values = bool(entries) and (kind != "donut" or any(value > 0 for value in values))
            chart_type = {
                "bars": XL_CHART_TYPE.BAR_CLUSTERED,
                "columns": XL_CHART_TYPE.COLUMN_CLUSTERED,
                "donut": XL_CHART_TYPE.DOUGHNUT,
            }[kind]

        if not has_values:
            empty_box = slide.shapes.add_textbox(
                Inches(left + 0.45),
                Inches(top + 0.95),
                Inches(width - 0.9),
                Inches(height - 1.2),
            )
            set_text(
                empty_box.text_frame,
                "No live values were returned for this diagram. Review Operational Notes for scope or access details.",
                size=10,
                color=palette["slate"],
                align=PP_ALIGN.CENTER,
            )
            return

        chart = slide.shapes.add_chart(
            chart_type,
            Inches(left + 0.12),
            Inches(top + 0.68),
            Inches(width - 0.24),
            Inches(height - 0.78),
            chart_data,
        ).chart
        chart.chart_style = 10
        chart.has_title = False
        chart.has_legend = kind in {"donut", "share"}
        if chart.has_legend:
            chart.legend.position = XL_LEGEND_POSITION.BOTTOM
            chart.legend.include_in_layout = False
            chart.legend.font.size = Pt(7)

        plot = chart.plots[0]
        value_format = str(panel.get("valueFormat", "number"))
        number_format = {
            "currency": "$#,##0.00",
            "days": '0" days"',
            "number": "#,##0.##",
            "percent": '0.0"%"',
        }.get(value_format, "#,##0.##")
        if kind == "share":
            number_format = "0%"
        if kind == "donut":
            plot.hole_size = 55
            plot.vary_by_categories = True
            plot.has_data_labels = True
            plot.data_labels.show_percentage = True
            plot.data_labels.show_category_name = False
            plot.data_labels.font.size = Pt(7)
        else:
            chart.category_axis.tick_labels.font.size = Pt(7)
            chart.value_axis.tick_labels.font.size = Pt(7)
            chart.value_axis.tick_labels.number_format = number_format
            chart.value_axis.has_major_gridlines = True
            plot.has_data_labels = True
            plot.data_labels.show_value = kind != "share"
            plot.data_labels.show_percentage = kind == "share"
            plot.data_labels.number_format = number_format
            plot.data_labels.font.size = Pt(7)
            if kind == "bars":
                chart.category_axis.reverse_order = True

    def add_diagram_slide(diagram: dict[str, object]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        section_slides["Report Diagram"] = len(presentation.slides)

        header = slide.shapes.add_textbox(Inches(0.5), Inches(0.18), Inches(8.4), Inches(0.4))
        set_text(
            header.text_frame,
            f"Live Report Diagram: {diagram['title']}",
            size=20,
            bold=True,
            color=palette["navy"],
        )
        freshness = slide.shapes.add_textbox(Inches(9.0), Inches(0.22), Inches(3.8), Inches(0.28))
        set_text(
            freshness.text_frame,
            f"Collected {report.generated_at_utc.strftime('%Y-%m-%d %H:%M UTC')} | {len(report.subscription_scope)} subscription(s)",
            size=8,
            color=palette["slate"],
            align=PP_ALIGN.RIGHT,
        )

        definition = slide.shapes.add_shape(
            MSO_SHAPE.RECTANGLE,
            Inches(0.5), Inches(0.68), Inches(12.3), Inches(0.72),
        )
        definition.fill.solid()
        definition.fill.fore_color.rgb = RGBColor(240, 253, 250)
        definition.line.color.rgb = palette["light"]
        scope_box = slide.shapes.add_textbox(Inches(0.72), Inches(0.78), Inches(5.7), Inches(0.5))
        set_text(scope_box.text_frame, f"Scope: {diagram['scope']}", size=9, color=palette["slate"])
        logic_box = slide.shapes.add_textbox(Inches(6.55), Inches(0.78), Inches(6.0), Inches(0.5))
        set_text(logic_box.text_frame, f"Logic: {diagram['logic']}", size=9, color=palette["slate"])

        panels = list(diagram["panels"])
        panel_width = 6.02
        panel_height = 2.55
        positions = [
            (0.5, 1.55),
            (6.78, 1.55),
            (0.5, 4.28),
            (6.78, 4.28),
        ]
        for panel, (left, top) in zip(panels[:4], positions, strict=False):
            add_chart_panel(slide, panel, left, top, panel_width, panel_height)

    def add_contents_slide():
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        header = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.5))
        set_text(header.text_frame, "Report Contents", size=22, bold=True, color=palette["navy"])
        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.6))
        text_frame = body.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        return text_frame

    def render_contents(text_frame, topics: list[tuple[str, str]]) -> None:
        for index, (title, context) in enumerate(topics):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            slide_number = section_slides.get(title)
            page_reference = f"Slide {slide_number}" if slide_number else "Section"
            paragraph.text = f"{index + 1}. {title} - {context}  |  {page_reference}"
            paragraph.space_after = Pt(5)
            paragraph.line_spacing = 1.0
            if paragraph.runs:
                paragraph.runs[0].font.size = Pt(13)
                paragraph.runs[0].font.color.rgb = palette["slate"]

    def add_bullet_slide(title: str, bullets: list[str]) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        header = slide.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.0), Inches(0.5))
        set_text(header.text_frame, title, size=22, bold=True, color=palette["navy"])

        body = slide.shapes.add_textbox(Inches(0.7), Inches(1.0), Inches(12.0), Inches(5.5))
        text_frame = body.text_frame
        text_frame.clear()
        text_frame.word_wrap = True
        text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
        for index, line in enumerate(bullets):
            paragraph = text_frame.paragraphs[0] if index == 0 else text_frame.add_paragraph()
            paragraph.text = line
            paragraph.level = 0
            paragraph.space_after = Pt(6)
            paragraph.line_spacing = 1.0
            if paragraph.runs:
                run = paragraph.runs[0]
                run.font.size = Pt(14)
                run.font.color.rgb = palette["slate"]

    def add_section_divider(title: str, subtitle: str) -> None:
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        section_slides[title] = len(presentation.slides)
        banner = slide.shapes.add_shape(1, Inches(0), Inches(0), presentation.slide_width, presentation.slide_height)
        banner.fill.solid()
        banner.fill.fore_color.rgb = palette["navy"]
        banner.line.fill.background()

        title_box = slide.shapes.add_textbox(Inches(0.75), Inches(2.0), Inches(11.8), Inches(0.7))
        set_text(title_box.text_frame, title, size=24, bold=True, color=palette["white"])

        subtitle_box = slide.shapes.add_textbox(Inches(0.75), Inches(2.8), Inches(11.8), Inches(0.5))
        set_text(subtitle_box.text_frame, subtitle, size=12, bold=False, color=palette["light"])

    def add_footers() -> None:
        total_slides = len(presentation.slides)
        for index, slide in enumerate(presentation.slides, start=1):
            footer = slide.shapes.add_textbox(Inches(0.5), Inches(7.0), Inches(12.0), Inches(0.22))
            footer_frame = footer.text_frame
            footer_frame.clear()
            paragraph = footer_frame.paragraphs[0]
            paragraph.text = f"FinOps CCoE Live Reporting | Slide {index} of {total_slides}"
            paragraph.alignment = PP_ALIGN.RIGHT
            if paragraph.runs:
                run = paragraph.runs[0]
                run.font.size = Pt(8)
                run.font.color.rgb = palette["slate"]

    def add_table_slide(title: str, rows: list[dict[str, object]], preferred_columns: list[str], rows_per_slide: int = 12) -> None:
        if not rows:
            return

        columns = [col for col in preferred_columns if any(col in row for row in rows)]
        for row in rows:
            columns.extend(col for col in row if col not in columns)

        index_candidates = [
            "serviceScope",
            "subscriptionId",
            "resourceName",
            "category",
            "source",
            "serviceName",
            "scope",
            "resourceId",
        ]
        index_columns = [col for col in index_candidates if col in columns][:3]
        detail_columns = [col for col in columns if col not in index_columns]
        fixed_columns = ["__rowNumber", *index_columns]
        detail_chunk_size = max(1, 6 - len(fixed_columns))
        column_groups = [
            [*fixed_columns, *detail_columns[start:start + detail_chunk_size]]
            for start in range(0, len(detail_columns), detail_chunk_size)
        ]
        if not column_groups:
            column_groups = [fixed_columns]

        for group_index, group_columns in enumerate(column_groups, start=1):
            group_rows_per_slide = rows_per_slide
            if len(group_columns) >= 4:
                group_rows_per_slide = min(rows_per_slide, 8)
            elif len(group_columns) >= 3:
                group_rows_per_slide = min(rows_per_slide, 10)

            for start in range(0, len(rows), group_rows_per_slide):
                chunk = rows[start:start + group_rows_per_slide]
                slide = presentation.slides.add_slide(presentation.slide_layouts[6])
                header = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(12.0), Inches(0.45))
                suffix = f" ({start + 1}-{start + len(chunk)} of {len(rows)})"
                if len(column_groups) > 1:
                    column_label = f" [Part {group_index}/{len(column_groups)}]"
                else:
                    column_label = ""
                set_text(header.text_frame, f"{title}{column_label}{suffix}", size=18, bold=True, color=palette["navy"])

                max_table_width = 12.5
                preferred_widths = []
                for col in group_columns:
                    values = [str(row.get(col, "")) for row in chunk]
                    header_text = "Row" if col == "__rowNumber" else col
                    longest = max([len(header_text), *(min(len(value), 72) for value in values)])
                    if col == "__rowNumber":
                        preferred_widths.append(0.55)
                    elif col == "resourceId":
                        preferred_widths.append(min(7.5, max(3.5, longest * 0.085)))
                    elif col in {"action", "evidence", "endpoint", "warning"}:
                        preferred_widths.append(min(6.0, max(2.4, longest * 0.085)))
                    elif col in {"resourceName", "serviceName", "serviceScope", "resourceType", "resourceGroup", "category", "scope"}:
                        preferred_widths.append(min(4.0, max(1.5, longest * 0.085)))
                    else:
                        preferred_widths.append(min(2.8, max(1.15, longest * 0.085)))

                preferred_total = sum(preferred_widths)
                if preferred_total > max_table_width:
                    scale = max_table_width / preferred_total
                    widths = [width * scale for width in preferred_widths]
                else:
                    widths = preferred_widths
                table_width = sum(widths)

                header_height = 0.38
                data_row_height = min(0.62, max(0.36, 5.7 / max(1, len(chunk))))
                table_height = header_height + (data_row_height * len(chunk))
                table_shape = slide.shapes.add_table(
                    len(chunk) + 1,
                    len(group_columns),
                    Inches(0.4),
                    Inches(0.8),
                    Inches(table_width),
                    Inches(table_height),
                )
                table = table_shape.table

                for idx, width in enumerate(widths):
                    table.columns[idx].width = Inches(width)
                table.rows[0].height = Inches(header_height)
                for row_index in range(1, len(table.rows)):
                    table.rows[row_index].height = Inches(data_row_height)

                for col_idx, col in enumerate(group_columns):
                    cell = table.cell(0, col_idx)
                    cell.text = "Row" if col == "__rowNumber" else col
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = palette["teal"]
                    cell.margin_left = Inches(0.08)
                    cell.margin_right = Inches(0.08)
                    cell.margin_top = Inches(0.04)
                    cell.margin_bottom = Inches(0.04)
                    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                    cell.text_frame.word_wrap = True
                    cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                    if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
                        run = cell.text_frame.paragraphs[0].runs[0]
                        run.font.bold = True
                        run.font.size = Pt(9)
                        run.font.color.rgb = palette["white"]

                for row_idx, row in enumerate(chunk, start=1):
                    row_fill = palette["white"] if row_idx % 2 else palette["pale"]
                    for col_idx, col in enumerate(group_columns):
                        value = start + row_idx if col == "__rowNumber" else _pptx_safe_text(row.get(col, ""), 28)
                        cell = table.cell(row_idx, col_idx)
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = row_fill
                        cell.text = str(value)
                        cell.margin_left = Inches(0.08)
                        cell.margin_right = Inches(0.08)
                        cell.margin_top = Inches(0.04)
                        cell.margin_bottom = Inches(0.04)
                        cell.vertical_anchor = MSO_ANCHOR.TOP
                        cell.text_frame.word_wrap = True
                        cell.text_frame.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
                        if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
                            run = cell.text_frame.paragraphs[0].runs[0]
                            run.font.size = Pt(8 if len(group_columns) >= 4 else 9)
                            run.font.color.rgb = palette["slate"]

    report_topics = [
        ("Executive Overview", "summary, priorities, action plan, and decisions"),
        ("Report Results", "live records for the selected report topic"),
    ]
    diagram = build_report_diagram(report)
    if diagram:
        report_topics.insert(1, ("Report Diagram", "scope, logic, and domain-specific live diagrams"))
    if report.recommendations:
        report_topics.insert(1, ("Recommendations", "impact, confidence, scope, actions, and evidence"))
    if report.ai_token_summary or report.ai_token_tracking:
        report_topics.insert(-1, ("AI Usage", "token and cost summaries"))
    if report.source_audit or report.enrichment_audit:
        report_topics.append(("Data Provenance", "source and enrichment audit context"))
    if report.warnings:
        report_topics.append(("Operational Notes", "warnings and partial-access conditions"))

    add_title_slide()
    contents_frame = add_contents_slide()
    add_section_divider("Executive Overview", "Current-run summary, priorities, action plan, and decisions.")
    add_summary_slide()
    if diagram:
        add_diagram_slide(diagram)
    add_bullet_slide(
        "Executive Priorities",
        [f"{title}: {detail}" for title, detail in brief["priorities"]],
    )
    add_bullet_slide(
        "90-Day Action Plan",
        [f"{period}: {action}" for period, action in brief["actions"]],
    )
    add_bullet_slide(
        "Executive Decisions",
        [str(decision) for decision in brief["executive_decisions"]],
    )

    if report.recommendations:
        add_section_divider("Recommendations", "Impact, confidence, scope, actions, and supporting evidence.")
        add_table_slide(
            "Recommendations",
            [
                {
                    "category": rec.category,
                    "impact": rec.impact_estimate,
                    "confidence": rec.confidence,
                    "scope": rec.affected_scope,
                    "action": rec.suggested_action,
                    "evidence": rec.evidence,
                }
                for rec in report.recommendations
            ],
            ["category", "impact", "confidence", "scope", "action", "evidence"],
            rows_per_slide=10,
        )

    if report.ai_token_summary or report.ai_token_tracking:
        add_section_divider("AI Usage", "Token and cost summaries for Azure AI services when available.")
    add_table_slide(
        "AI Token Summary",
        report.ai_token_summary,
        [
            "serviceScope",
            "records",
            "totalUsageQuantity",
            "averageUsageQuantity",
            "totalEstimatedTokenUnits",
            "averageEstimatedTokenUnits",
            "totalCost",
            "averageCost",
        ],
    )

    add_table_slide(
        "AI Token Tracking",
        report.ai_token_tracking,
        [
            "subscriptionId",
            "serviceName",
            "meterName",
            "usageQuantity",
            "estimatedTokenUnits",
            "cost",
            "resourceId",
        ],
    )

    add_section_divider("Report Results", "Live records for the selected report topic and subscription scope.")
    add_table_slide(
        "Report Results",
        report.rows,
        [
            "subscriptionId",
            "resourceName",
            "serviceName",
            "resourceType",
            "resourceGroup",
            "region",
            "cost",
            "usageQuantity",
            "estimatedTokenUnits",
            "resourceId",
        ],
    )

    if report.source_audit or report.enrichment_audit:
        add_section_divider("Data Provenance", "Source calls and enrichment outcomes supporting this report.")
    add_table_slide(
        "Source Audit",
        [
            {
                "source": item.source,
                "endpoint": item.endpoint,
                "durationMs": item.duration_ms,
                "rowCount": item.row_count,
                "executedAtUtc": item.executed_at_utc.isoformat(),
            }
            for item in report.source_audit
        ],
        ["source", "durationMs", "rowCount", "executedAtUtc", "endpoint"],
        rows_per_slide=10,
    )

    add_table_slide(
        "Enrichment Audit",
        report.enrichment_audit,
        [
            "resolutionStatus",
            "resolvedResourceName",
            "resolvedResourceType",
            "resolvedResourceGroup",
            "resolvedSubscriptionId",
            "sourceResourceId",
        ],
    )

    if report.warnings:
        add_section_divider("Operational Notes", "Warnings and partial-access notes.")
        add_table_slide(
            "Warnings",
            [{"warning": warning} for warning in report.warnings],
            ["warning"],
            rows_per_slide=15,
        )

    render_contents(contents_frame, report_topics)
    add_footers()

    presentation.save(str(output))
    return output
